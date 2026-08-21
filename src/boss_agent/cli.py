from __future__ import annotations

import argparse
import csv
import json
import re
import time
import random
from datetime import date
from datetime import datetime
from pathlib import Path
from typing import Any

from boss_agent import __version__
from boss_agent.capability_agent import analyze_operations_capability
from boss_agent.chrome_debug import ChromeDebugClient
from boss_agent.chrome_debug import build_chrome_launch_command
from boss_agent.chrome_debug import classify_page
from boss_agent.chrome_debug import launch_chrome
from boss_agent.drafting import draft_reply
from boss_agent.recommendation import score_recommendation
from boss_agent.resume_agent import analyze_resume_against_criteria
from boss_agent.scoring import score_conversation
from boss_agent.knowledge_base import answer_question_from_knowledge_detailed
from boss_agent.knowledge_base import load_knowledge_base
from boss_agent.knowledge_base import normalize_text
from boss_agent.job_agent import load_local_llm_env


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="boss-agent")
    subparsers = parser.add_subparsers(dest="command", required=True)

    subparsers.add_parser("health")
    capture_parser = subparsers.add_parser("capture")
    capture_subparsers = capture_parser.add_subparsers(dest="capture_command", required=True)
    score_parser = subparsers.add_parser("score")
    score_subparsers = score_parser.add_subparsers(dest="score_command", required=True)
    draft_parser = subparsers.add_parser("draft")
    draft_subparsers = draft_parser.add_subparsers(dest="draft_command", required=True)
    action_parser = subparsers.add_parser("action")
    action_subparsers = action_parser.add_subparsers(dest="action_command", required=True)

    conversation_parser = capture_subparsers.add_parser("conversation")
    conversation_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    conversation_parser.add_argument("--mock-dir")
    conversation_parser.add_argument("--url-contains", default="/web/geek/chat")

    structured_parser = capture_subparsers.add_parser("conversation-structured")
    structured_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    structured_parser.add_argument("--mock-dir")
    structured_parser.add_argument("--url-contains", default="/web/geek/chat")

    attachment_link_parser = capture_subparsers.add_parser("attachment-link")
    attachment_link_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    attachment_link_parser.add_argument("--mock-dir")
    attachment_link_parser.add_argument("--url-contains", default="/web/chat/index")

    inbox_unread_parser = capture_subparsers.add_parser("inbox-unread")
    inbox_unread_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    inbox_unread_parser.add_argument("--mock-dir")
    inbox_unread_parser.add_argument("--url-contains", default="/web/chat/index")

    job_filter_options_parser = capture_subparsers.add_parser("job-filter-options")
    job_filter_options_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    job_filter_options_parser.add_argument("--mock-dir")
    job_filter_options_parser.add_argument("--url-contains", default="/web/chat/index")

    recommend_candidates_parser = capture_subparsers.add_parser("recommend-candidates")
    recommend_candidates_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    recommend_candidates_parser.add_argument("--mock-dir")
    recommend_candidates_parser.add_argument("--url-contains", default="/web/chat/recommend")
    recommend_candidates_parser.add_argument("--max-count", type=int, default=20)

    recommend_resume_cards_parser = capture_subparsers.add_parser("recommend-resume-cards")
    recommend_resume_cards_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    recommend_resume_cards_parser.add_argument("--mock-dir")
    recommend_resume_cards_parser.add_argument("--url-contains", default="/web/chat/recommend")
    recommend_resume_cards_parser.add_argument("--max-scrolls", type=int, default=12)

    score_conversation_parser = score_subparsers.add_parser("conversation")
    score_conversation_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    score_conversation_parser.add_argument("--mock-dir")
    score_conversation_parser.add_argument("--url-contains", default="/web/geek/chat")

    score_recommendation_parser = score_subparsers.add_parser("recommendation")
    score_recommendation_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    score_recommendation_parser.add_argument("--mock-dir")
    score_recommendation_parser.add_argument("--url-contains", default="/web/chat/recommend")
    score_recommendation_parser.add_argument("--job-title", required=True)
    score_recommendation_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))
    score_recommendation_parser.add_argument("--min-score", type=int, default=70)

    draft_reply_parser = draft_subparsers.add_parser("reply")
    draft_reply_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    draft_reply_parser.add_argument("--mock-dir")
    draft_reply_parser.add_argument("--url-contains", default="/web/geek/chat")
    draft_reply_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))

    draft_fill_parser = draft_subparsers.add_parser("fill")
    draft_fill_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    draft_fill_parser.add_argument("--mock-dir")
    draft_fill_parser.add_argument("--url-contains", default="/web/geek/chat")
    draft_fill_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))

    resume_consent_parser = action_subparsers.add_parser("resume-consent")
    resume_consent_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    resume_consent_parser.add_argument("--mock-dir")
    resume_consent_parser.add_argument("--url-contains", default="/web/chat/index")

    download_resume_parser = action_subparsers.add_parser("download-resume")
    download_resume_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    download_resume_parser.add_argument("--mock-dir")
    download_resume_parser.add_argument("--url-contains", default="/web/chat/index")
    download_resume_parser.add_argument("--downloads-dir", default=str(Path.home() / "Downloads"))

    open_conversation_parser = action_subparsers.add_parser("open-conversation")
    open_conversation_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    open_conversation_parser.add_argument("--mock-dir")
    open_conversation_parser.add_argument("--url-contains", default="/web/chat/index")
    open_conversation_parser.add_argument("--data-id")
    open_conversation_parser.add_argument("--candidate-name")

    prepare_reply_parser = action_subparsers.add_parser("prepare-reply")
    prepare_reply_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    prepare_reply_parser.add_argument("--mock-dir")
    prepare_reply_parser.add_argument("--url-contains", default="/web/chat/index")
    prepare_reply_parser.add_argument("--downloads-dir", default=str(Path.home() / "Downloads"))
    prepare_reply_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))

    close_attachment_preview_parser = action_subparsers.add_parser("close-attachment-preview")
    close_attachment_preview_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    close_attachment_preview_parser.add_argument("--mock-dir")
    close_attachment_preview_parser.add_argument("--url-contains", default="/web/chat/index")

    switch_job_filter_parser = action_subparsers.add_parser("switch-job-filter")
    switch_job_filter_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    switch_job_filter_parser.add_argument("--mock-dir")
    switch_job_filter_parser.add_argument("--url-contains", default="/web/chat/index")
    switch_job_filter_parser.add_argument("--job-text", required=True)

    traverse_job_search_parser = action_subparsers.add_parser("traverse-job-search")
    traverse_job_search_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    traverse_job_search_parser.add_argument("--mock-dir")
    traverse_job_search_parser.add_argument("--url-contains", default="zhipin.com")
    traverse_job_search_parser.add_argument("--max-count", type=int, default=20)
    traverse_job_search_parser.add_argument("--max-scrolls", type=int, default=20)
    traverse_job_search_parser.add_argument("--wait-timeout-seconds", type=float, default=15.0)
    traverse_job_search_parser.add_argument("--send", action="store_true")
    traverse_job_search_parser.add_argument("--max-communications", type=int)
    traverse_job_search_parser.add_argument("--model")
    traverse_job_search_parser.add_argument("--operation-delay-seconds", type=float, default=1.5)
    traverse_job_search_parser.add_argument("--expectation")

    process_unread_parser = action_subparsers.add_parser("process-unread")
    process_unread_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    process_unread_parser.add_argument("--mock-dir")
    process_unread_parser.add_argument("--url-contains", default="/web/chat/index")
    process_unread_parser.add_argument("--downloads-dir", default=str(Path.home() / "Downloads"))
    process_unread_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))
    process_unread_parser.add_argument("--max-count", type=int, default=10)
    process_unread_parser.add_argument("--delay-seconds", type=float, default=3.0)
    process_unread_parser.add_argument("--send", action="store_true")
    process_unread_parser.add_argument("--fallback-today", action="store_true")

    send_unread_message_parser = action_subparsers.add_parser("send-unread-message")
    send_unread_message_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    send_unread_message_parser.add_argument("--mock-dir")
    send_unread_message_parser.add_argument("--url-contains", default="/web/chat/index")
    send_unread_message_parser.add_argument("--job-text", required=True)
    send_unread_message_parser.add_argument("--message", required=True)
    send_unread_message_parser.add_argument("--max-count", type=int, default=50)
    send_unread_message_parser.add_argument("--inbox-scrolls", type=int, default=20)
    send_unread_message_parser.add_argument("--delay-seconds", type=float, default=1.0)
    send_unread_message_parser.add_argument("--dry-run", action="store_true")

    request_resume_parser = action_subparsers.add_parser("request-resume-new-greetings")
    request_resume_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    request_resume_parser.add_argument("--mock-dir")
    request_resume_parser.add_argument("--url-contains", default="/web/chat/index")
    request_resume_parser.add_argument("--job-text", required=True)
    request_resume_parser.add_argument("--message", default="您好，方便发一份简历吗？我这边先看一下。")
    request_resume_parser.add_argument("--criteria", default="")
    request_resume_parser.add_argument("--max-count", type=int, default=50)
    request_resume_parser.add_argument("--inbox-scrolls", type=int, default=20)
    request_resume_parser.add_argument("--wait-after-unread", type=float, default=3.0)
    request_resume_parser.add_argument("--operation-delay-seconds", type=float, default=1.0)
    request_resume_parser.add_argument("--send", action="store_true")

    knowledge_reply_parser = action_subparsers.add_parser("reply-unread-with-knowledge")
    knowledge_reply_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    knowledge_reply_parser.add_argument("--mock-dir")
    knowledge_reply_parser.add_argument("--url-contains", default="/web/chat/index")
    knowledge_reply_parser.add_argument("--job-text", required=True)
    knowledge_reply_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))
    knowledge_reply_parser.add_argument("--knowledge-file")
    knowledge_reply_parser.add_argument("--max-count", type=int, default=50)
    knowledge_reply_parser.add_argument("--inbox-scrolls", type=int, default=20)
    knowledge_reply_parser.add_argument("--operation-delay-seconds", type=float, default=2.0)
    knowledge_reply_parser.add_argument("--send", action="store_true")

    process_all_today_parser = action_subparsers.add_parser("process-all-today")
    process_all_today_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    process_all_today_parser.add_argument("--mock-dir")
    process_all_today_parser.add_argument("--url-contains", default="/web/chat/index")
    process_all_today_parser.add_argument("--downloads-dir", default=str(Path.home() / "Downloads"))
    process_all_today_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))
    process_all_today_parser.add_argument("--max-count", type=int, default=10)
    process_all_today_parser.add_argument("--min-delay-seconds", type=float, default=30.0)
    process_all_today_parser.add_argument("--max-delay-seconds", type=float, default=60.0)
    process_all_today_parser.add_argument("--send", action="store_true")

    search_recommendations_parser = action_subparsers.add_parser("search-recommendations")
    search_recommendations_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    search_recommendations_parser.add_argument("--mock-dir")
    search_recommendations_parser.add_argument("--url-contains", default="/web/chat/recommend")
    search_recommendations_parser.add_argument("--job-title", required=True)
    search_recommendations_parser.add_argument("--city", default="")
    search_recommendations_parser.add_argument("--filter", action="append", default=[])

    process_recommendations_parser = action_subparsers.add_parser("process-recommendations")
    process_recommendations_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    process_recommendations_parser.add_argument("--mock-dir")
    process_recommendations_parser.add_argument("--url-contains", default="/web/chat/recommend")
    process_recommendations_parser.add_argument("--job-title", required=True)
    process_recommendations_parser.add_argument("--city", default="")
    process_recommendations_parser.add_argument("--filter", action="append", default=[])
    process_recommendations_parser.add_argument("--knowledge-dir", default=str((Path.cwd() / "knowledge").resolve()))
    process_recommendations_parser.add_argument("--max-count", type=int, default=20)
    process_recommendations_parser.add_argument("--min-score", type=int, default=70)
    process_recommendations_parser.add_argument("--send", action="store_true")

    filter_recommendations_parser = action_subparsers.add_parser("filter-recommendations")
    filter_recommendations_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    filter_recommendations_parser.add_argument("--mock-dir")
    filter_recommendations_parser.add_argument("--url-contains", default="/web/chat/recommend")
    filter_recommendations_parser.add_argument("--criteria", required=True)
    filter_recommendations_parser.add_argument("--max-scrolls", type=int, default=12)
    filter_recommendations_parser.add_argument("--model")

    greet_recommendations_parser = action_subparsers.add_parser("greet-recommendations")
    greet_recommendations_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    greet_recommendations_parser.add_argument("--mock-dir")
    greet_recommendations_parser.add_argument("--url-contains", default="/web/chat/recommend")
    greet_recommendations_parser.add_argument("--name", action="append", default=[])

    screen_and_greet_parser = action_subparsers.add_parser("screen-and-greet-recommendations")
    screen_and_greet_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    screen_and_greet_parser.add_argument("--mock-dir")
    screen_and_greet_parser.add_argument("--url-contains", default="/web/chat/recommend")
    screen_and_greet_parser.add_argument("--criteria", required=True)
    screen_and_greet_parser.add_argument("--max-scrolls", type=int, default=12)
    screen_and_greet_parser.add_argument("--model")
    screen_and_greet_parser.add_argument("--max-greetings", type=int, default=20)
    screen_and_greet_parser.add_argument("--dry-run", action="store_true")

    capability_greet_parser = action_subparsers.add_parser("screen-and-greet-capability")
    capability_greet_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    capability_greet_parser.add_argument("--mock-dir")
    capability_greet_parser.add_argument("--url-contains", default="/web/chat/recommend")
    capability_greet_parser.add_argument("--criteria", required=True)
    capability_greet_parser.add_argument("--max-scrolls", type=int, default=18)
    capability_greet_parser.add_argument("--max-greetings", type=int, default=30)
    capability_greet_parser.add_argument("--report-path")
    capability_greet_parser.add_argument("--dry-run", action="store_true")

    detail_greet_parser = action_subparsers.add_parser("screen-recommend-detail")
    detail_greet_parser.add_argument("--config")
    detail_greet_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    detail_greet_parser.add_argument("--mock-dir")
    detail_greet_parser.add_argument("--url-contains", default="/web/chat/recommend")
    detail_greet_parser.add_argument("--job-title", default="")
    detail_greet_parser.add_argument("--city", default="")
    detail_greet_parser.add_argument("--filter", action="append", default=[])
    detail_greet_parser.add_argument("--criteria")
    detail_greet_parser.add_argument("--model")
    detail_greet_parser.add_argument("--max-greetings", type=int)
    detail_greet_parser.add_argument("--max-checks", type=int, default=200)
    detail_greet_parser.add_argument("--target-resumes", type=int)
    detail_greet_parser.add_argument("--post-filter-wait", type=float, default=5.0)
    detail_greet_parser.add_argument("--operation-delay-seconds", type=float)
    detail_greet_parser.add_argument("--max-filter-retries", type=int)
    detail_greet_parser.add_argument("--report-path", default=str(Path.cwd() / "reports" / "recommend_detail_greeted.xlsx"))
    detail_greet_parser.add_argument("--dry-run", action="store_true")

    debug_parser = subparsers.add_parser("debug")
    debug_subparsers = debug_parser.add_subparsers(dest="debug_command", required=True)

    version_parser = debug_subparsers.add_parser("version")
    version_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    version_parser.add_argument("--mock-file")

    pages_parser = debug_subparsers.add_parser("pages")
    pages_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    pages_parser.add_argument("--mock-dir")

    detect_parser = debug_subparsers.add_parser("detect-page")
    detect_parser.add_argument("--endpoint", default="http://127.0.0.1:9222")
    detect_parser.add_argument("--mock-dir")

    launch_parser = debug_subparsers.add_parser("launch-chrome")
    launch_parser.add_argument(
        "--chrome-path",
        default=r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    )
    launch_parser.add_argument(
        "--user-data-dir",
        default=str((Path.cwd() / ".chrome-profile").resolve()),
    )
    launch_parser.add_argument("--port", type=int, default=9222)
    launch_parser.add_argument("--dry-run", action="store_true")

    open_boss_parser = debug_subparsers.add_parser("open-boss")
    open_boss_parser.add_argument(
        "--chrome-path",
        default=r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    )
    open_boss_parser.add_argument(
        "--user-data-dir",
        default=str((Path.cwd() / ".chrome-profile").resolve()),
    )
    open_boss_parser.add_argument("--port", type=int, default=9222)
    open_boss_parser.add_argument(
        "--page",
        choices=["recommend", "chat"],
        default="chat",
    )
    open_boss_parser.add_argument("--url")
    open_boss_parser.add_argument("--dry-run", action="store_true")

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if args.command == "health":
        return _print_json(
            {
                "status": "ok",
                "service": "boss-agent",
                "version": __version__,
            }
        )

    if args.command == "capture":
        if args.capture_command == "conversation":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            captured = client.capture_visible_text(url_contains=args.url_contains)
            page = classify_page({"title": captured["title"], "url": captured["url"], "type": "page"})
            lines = [line.strip() for line in captured["text"].splitlines() if line.strip()]
            return _print_json(
                {
                    "title": captured["title"],
                    "url": captured["url"],
                    "platform": page["platform"],
                    "pageType": page["pageType"],
                    "text": captured["text"],
                    "preview": lines[:8],
                }
            )
        if args.capture_command == "conversation-structured":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            captured = client.capture_conversation(url_contains=args.url_contains)
            return _print_json(captured)
        if args.capture_command == "attachment-link":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(client.capture_attachment_preview_link(url_contains=args.url_contains))
        if args.capture_command == "inbox-unread":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(client.capture_unread_inbox(url_contains=args.url_contains))
        if args.capture_command == "job-filter-options":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(client.capture_job_filter_options(url_contains=args.url_contains))
        if args.capture_command == "recommend-candidates":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.capture_recommend_candidates(
                    max_count=args.max_count,
                    url_contains=args.url_contains,
                )
            )
        if args.capture_command == "recommend-resume-cards":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.capture_recommend_resume_cards(
                    max_scrolls=args.max_scrolls,
                    url_contains=args.url_contains,
                )
            )

    if args.command == "score":
        if args.score_command == "conversation":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            if fixture_dir:
                captured = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir).capture_conversation(
                    url_contains=args.url_contains
                )
            else:
                captured = ChromeDebugClient(endpoint=args.endpoint).capture_conversation(
                    url_contains=args.url_contains
                )
            return _print_json(score_conversation(captured))
        if args.score_command == "recommendation":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            resume = client.capture_candidate_resume(url_contains=args.url_contains)
            return _print_json(
                score_recommendation(
                    resume,
                    job_title=args.job_title,
                    knowledge_dir=Path(args.knowledge_dir),
                    min_score=args.min_score,
                )
            )

    if args.command == "draft":
        if args.draft_command == "reply":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            if fixture_dir:
                captured = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir).capture_conversation(
                    url_contains=args.url_contains
                )
            else:
                captured = ChromeDebugClient(endpoint=args.endpoint).capture_conversation(
                    url_contains=args.url_contains
                )
            score = score_conversation(captured)
            return _print_json(draft_reply(captured, score, knowledge_dir=Path(args.knowledge_dir)))
        if args.draft_command == "fill":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            captured = client.capture_conversation(url_contains=args.url_contains)
            score = score_conversation(captured)
            draft = draft_reply(captured, score, knowledge_dir=Path(args.knowledge_dir))
            return _print_json(client.fill_chat_input(draft["body"], url_contains=args.url_contains))

    if args.command == "action":
        if args.action_command == "resume-consent":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(client.click_resume_consent(url_contains=args.url_contains))
        if args.action_command == "download-resume":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.download_resume(
                    url_contains=args.url_contains,
                    downloads_dir=Path(args.downloads_dir),
                )
            )
        if args.action_command == "open-conversation":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.open_conversation(
                    data_id=args.data_id,
                    candidate_name=args.candidate_name,
                    url_contains=args.url_contains,
                )
            )
        if args.action_command == "prepare-reply":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                _prepare_reply(
                    client=client,
                    url_contains=args.url_contains,
                    downloads_dir=Path(args.downloads_dir),
                    knowledge_dir=Path(args.knowledge_dir),
                )
            )
        if args.action_command == "close-attachment-preview":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(client.close_attachment_preview(url_contains=args.url_contains))
        if args.action_command == "switch-job-filter":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.switch_job_filter(
                    job_text=args.job_text,
                    url_contains=args.url_contains,
                )
            )
        if args.action_command == "traverse-job-search":
            load_local_llm_env()
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.traverse_job_search(
                    max_count=args.max_count,
                    max_scrolls=args.max_scrolls,
                    wait_timeout_seconds=args.wait_timeout_seconds,
                    url_contains=args.url_contains,
                    send=args.send,
                    max_communications=args.max_communications,
                    model=args.model,
                    operation_delay_seconds=args.operation_delay_seconds,
                    expectation=args.expectation,
                )
            )
        if args.action_command == "process-unread":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            switch_result = client.switch_message_filter(label="未读", url_contains=args.url_contains)
            inbox = client.capture_unread_inbox(url_contains=args.url_contains)
            conversations = inbox.get("conversations", [])[: args.max_count]
            processed: list[dict[str, Any]] = []
            processed_ids: set[str] = set()

            for item in conversations:
                processed.append(
                    _process_conversation_item(
                        client=client,
                        item=item,
                        url_contains=args.url_contains,
                        downloads_dir=Path(args.downloads_dir),
                        knowledge_dir=Path(args.knowledge_dir),
                        send=args.send,
                        delay_seconds=args.delay_seconds,
                    )
                )
                if item.get("dataId"):
                    processed_ids.add(item["dataId"])

            fallback = None
            if args.fallback_today:
                fallback_switch = client.switch_message_filter(label="全部", url_contains=args.url_contains)
                fallback_inbox = client.capture_inbox(url_contains=args.url_contains, filter_label="全部")
                fallback_candidates = [
                    item for item in fallback_inbox.get("conversations", [])
                    if _is_today_time(item.get("time", "")) and item.get("dataId") not in processed_ids
                ][: args.max_count]
                fallback_processed: list[dict[str, Any]] = []
                for item in fallback_candidates:
                    fallback_processed.append(
                        _process_conversation_item(
                            client=client,
                            item=item,
                            url_contains=args.url_contains,
                            downloads_dir=Path(args.downloads_dir),
                            knowledge_dir=Path(args.knowledge_dir),
                            send=args.send,
                            delay_seconds=args.delay_seconds,
                        )
                    )
                    if item.get("dataId"):
                        processed_ids.add(item["dataId"])
                fallback = {
                    "switchFilter": fallback_switch,
                    "inbox": fallback_inbox,
                    "processedCount": len(fallback_processed),
                    "processed": fallback_processed,
                }

            return _print_json(
                {
                    "switchFilter": switch_result,
                    "inbox": inbox,
                    "processedCount": len(processed),
                    "processed": processed,
                    "sent": args.send,
                    "fallbackToday": fallback,
                }
            )
        if args.action_command == "send-unread-message":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                _send_unread_message(
                    client=client,
                    url_contains=args.url_contains,
                    job_text=args.job_text,
                    message=args.message,
                    max_count=args.max_count,
                    inbox_scrolls=args.inbox_scrolls,
                    delay_seconds=args.delay_seconds,
                    dry_run=args.dry_run,
                )
            )
        if args.action_command == "request-resume-new-greetings":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                _request_resume_new_greetings(
                    client=client,
                    url_contains=args.url_contains,
                    job_text=args.job_text,
                    message=args.message,
                    max_count=args.max_count,
                    inbox_scrolls=args.inbox_scrolls,
                    wait_after_unread=args.wait_after_unread,
                    operation_delay_seconds=args.operation_delay_seconds,
                    send=args.send,
                    criteria=args.criteria,
                )
            )
        if args.action_command == "reply-unread-with-knowledge":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                _reply_unread_with_knowledge(
                    client=client,
                    url_contains=args.url_contains,
                    job_text=args.job_text,
                    knowledge_dir=Path(args.knowledge_dir),
                    knowledge_file=Path(args.knowledge_file) if args.knowledge_file else None,
                    max_count=args.max_count,
                    inbox_scrolls=args.inbox_scrolls,
                    operation_delay_seconds=args.operation_delay_seconds,
                    send=args.send,
                )
            )
        if args.action_command == "process-all-today":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            switch_result = client.switch_message_filter(label="全部", url_contains=args.url_contains)
            inbox = client.capture_inbox(url_contains=args.url_contains, filter_label="全部")
            conversations = [
                item for item in inbox.get("conversations", [])
                if _is_today_time(item.get("time", ""))
            ][: args.max_count]
            processed: list[dict[str, Any]] = []

            for item in conversations:
                processed.append(
                    _process_conversation_item(
                        client=client,
                        item=item,
                        url_contains=args.url_contains,
                        downloads_dir=Path(args.downloads_dir),
                        knowledge_dir=Path(args.knowledge_dir),
                        send=args.send,
                        delay_seconds=_resolve_delay_seconds(
                            min_delay_seconds=args.min_delay_seconds,
                            max_delay_seconds=args.max_delay_seconds,
                        ),
                    )
                )

            return _print_json(
                {
                    "switchFilter": switch_result,
                    "inbox": inbox,
                    "processedCount": len(processed),
                    "processed": processed,
                    "sent": args.send,
                }
            )
        if args.action_command == "search-recommendations":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            open_result = client.open_recommend_page(url_contains=args.url_contains)
            search_result = client.apply_recommend_filters(
                job_title=args.job_title,
                city=args.city,
                filters=_parse_filter_args(args.filter),
                url_contains=args.url_contains,
            )
            return _print_json(
                {
                    "openRecommendPage": open_result,
                    "search": search_result,
                }
            )
        if args.action_command == "process-recommendations":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            result = _process_recommendations(
                client=client,
                url_contains=args.url_contains,
                job_title=args.job_title,
                city=args.city,
                filters=_parse_filter_args(args.filter),
                knowledge_dir=Path(args.knowledge_dir),
                max_count=args.max_count,
                min_score=args.min_score,
                send=args.send,
            )
            return _print_json(result)
        if args.action_command == "filter-recommendations":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            cards_payload = client.capture_recommend_resume_cards(
                max_scrolls=args.max_scrolls,
                url_contains=args.url_contains,
            )
            analyzed: list[dict[str, Any]] = []
            for card in cards_payload.get("cards", []):
                analysis = analyze_resume_against_criteria(
                    resume_text=str(card.get("rawText") or ""),
                    criteria=args.criteria,
                    model=args.model,
                )
                name = analysis.get("name") or card.get("name") or ""
                analyzed.append(
                    {
                        "name": name,
                        "summary": card.get("summary", ""),
                        "candidateSummary": analysis.get("candidateSummary", ""),
                        "meetsCriteria": analysis.get("meetsCriteria", False),
                        "educationLevel": analysis.get("educationLevel", ""),
                        "workYears": analysis.get("workYears"),
                        "agent": analysis.get("agent", ""),
                        "reasons": analysis.get("reasons", []),
                        "risks": analysis.get("risks", []),
                    }
                )
            matched = [item for item in analyzed if item.get("meetsCriteria")]
            matched_names = list(dict.fromkeys(str(item.get("name", "")) for item in matched if item.get("name")))
            return _print_json(
                {
                    "criteria": args.criteria,
                    "checkedCount": len(analyzed),
                    "matchedNames": matched_names,
                    "matched": matched,
                    "analyzed": analyzed,
                    "source": {
                        "url": cards_payload.get("url", ""),
                        "frameUrl": cards_payload.get("frameUrl", ""),
                        "cardCount": cards_payload.get("count", 0),
                    },
                }
            )
        if args.action_command == "greet-recommendations":
            if not args.name:
                parser.error("action greet-recommendations requires at least one --name")
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                client.greet_recommend_candidates_by_name(
                    names=args.name,
                    url_contains=args.url_contains,
                )
            )
        if args.action_command == "screen-and-greet-recommendations":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json(
                _screen_and_greet_recommendations(
                    client=client,
                    url_contains=args.url_contains,
                    criteria=args.criteria,
                    max_scrolls=args.max_scrolls,
                    model=args.model,
                    max_greetings=args.max_greetings,
                    dry_run=args.dry_run,
                )
            )
        if args.action_command == "screen-and-greet-capability":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            report_path = Path(args.report_path) if args.report_path else _default_capability_report_path()
            return _print_json(
                _screen_and_greet_capability(
                    client=client,
                    url_contains=args.url_contains,
                    criteria=args.criteria,
                    max_scrolls=args.max_scrolls,
                    max_greetings=args.max_greetings,
                    report_path=report_path,
                    dry_run=args.dry_run,
                )
            )
        if args.action_command == "screen-recommend-detail":
            config = _load_screen_recommend_detail_config(args)
            fixture_dir = Path(config["mock_dir"]).resolve() if config.get("mock_dir") else None
            client = ChromeDebugClient(endpoint=config["endpoint"], mock_dir=fixture_dir)
            result = _screen_recommend_detail(
                client=client,
                url_contains=config["url_contains"],
                job_title=config["job_title"],
                city=config["city"],
                filters=config["filters"],
                criteria=config["criteria"],
                model=config["model"],
                max_greetings=config["max_greetings"],
                max_checks=config["max_checks"],
                post_filter_wait_seconds=config["post_filter_wait_seconds"],
                operation_delay_seconds=config["operation_delay_seconds"],
                max_filter_retries=config["max_filter_retries"],
                report_path=Path(config["report_path"]),
                dry_run=config["dry_run"],
            )
            return _print_json(result)

    if args.command == "debug":
        if args.debug_command == "version":
            fixture_dir = None
            if args.mock_file:
                fixture_dir = Path(args.mock_file).resolve().parent
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            payload = client.version()
            return _print_json(
                {
                    "browser": payload.get("Browser"),
                    "protocolVersion": payload.get("Protocol-Version"),
                    "webSocketDebuggerUrl": payload.get("webSocketDebuggerUrl"),
                }
            )

        if args.debug_command == "pages":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            return _print_json({"pages": client.list_pages()})

        if args.debug_command == "detect-page":
            fixture_dir = Path(args.mock_dir).resolve() if args.mock_dir else None
            client = ChromeDebugClient(endpoint=args.endpoint, mock_dir=fixture_dir)
            pages = client.list_pages()
            classified_pages = [classify_page(page) for page in pages]
            matched_pages = [page for page in classified_pages if page["platform"] == "boss"]
            return _print_json(
                {
                    "pageCount": len(classified_pages),
                    "matchedCount": len(matched_pages),
                    "pages": matched_pages,
                }
            )

        if args.debug_command == "launch-chrome":
            chrome_path = Path(args.chrome_path)
            user_data_dir = Path(args.user_data_dir)
            if args.dry_run:
                command = build_chrome_launch_command(
                    chrome_path=chrome_path,
                    user_data_dir=user_data_dir,
                    port=args.port,
                )
            else:
                command = launch_chrome(
                    chrome_path=chrome_path,
                    user_data_dir=user_data_dir,
                    port=args.port,
                )
            return _print_json(
                {
                    "action": "launch",
                    "port": args.port,
                    "chromePath": str(chrome_path),
                    "userDataDir": str(user_data_dir),
                    "command": command,
                }
            )

        if args.debug_command == "open-boss":
            chrome_path = Path(args.chrome_path)
            user_data_dir = Path(args.user_data_dir)
            page_urls = {
                "recommend": "https://www.zhipin.com/web/chat/recommend",
                "chat": "https://www.zhipin.com/web/chat/index",
            }
            url = args.url or page_urls[args.page]
            if args.dry_run:
                command = build_chrome_launch_command(
                    chrome_path=chrome_path,
                    user_data_dir=user_data_dir,
                    port=args.port,
                    urls=[url],
                )
            else:
                command = launch_chrome(
                    chrome_path=chrome_path,
                    user_data_dir=user_data_dir,
                    port=args.port,
                    urls=[url],
                )
            return _print_json(
                {
                    "action": "open-boss",
                    "port": args.port,
                    "page": args.page,
                    "url": url,
                    "chromePath": str(chrome_path),
                    "userDataDir": str(user_data_dir),
                    "command": command,
                }
            )

    parser.error("Unsupported command")
    return 2


def _print_json(payload: dict[str, Any]) -> int:
    print(json.dumps(payload, ensure_ascii=False))
    return 0


def _prepare_reply(
    client: ChromeDebugClient,
    url_contains: str,
    downloads_dir: Path,
    knowledge_dir: Path,
) -> dict[str, Any]:
    conversation = client.capture_conversation(url_contains=url_contains)
    messages = conversation.get("messages", [])
    message_texts = [message.get("text", "") for message in messages]
    resume_prompt = client.detect_resume_consent_prompt(url_contains=url_contains)
    has_resume_request = any("对方想发送附件简历给您，您是否同意" in text for text in message_texts) or resume_prompt.get("visible", False)
    has_resume_attachment = any("点击预览附件简历" in text or "附件简历" in text for text in message_texts)

    resume_consent: dict[str, Any] | None = None
    attachment_preview: dict[str, Any] | None = None
    resume_download: dict[str, Any] | None = None

    if has_resume_request:
        resume_consent = client.click_resume_consent(url_contains=url_contains)
        time.sleep(1)

    if has_resume_request or has_resume_attachment:
        attachment_preview, resume_download = _download_resume_with_retry(
            client=client,
            url_contains=url_contains,
            downloads_dir=downloads_dir,
        )

    updated_conversation = client.capture_conversation(url_contains=url_contains)
    score = score_conversation(updated_conversation)
    draft = draft_reply(updated_conversation, score, knowledge_dir=knowledge_dir)
    draft_fill = client.fill_chat_input(draft["body"], url_contains=url_contains)

    return {
        "conversation": updated_conversation,
        "score": score,
        "draft": draft,
        "draftFill": draft_fill,
        "resumePrompt": resume_prompt,
        "resumeConsent": resume_consent,
        "attachmentPreview": attachment_preview,
        "resumeDownload": resume_download,
    }


def _load_screen_recommend_detail_config(args: argparse.Namespace) -> dict[str, Any]:
    file_config: dict[str, Any] = {}
    if args.config:
        config_path = Path(args.config).resolve()
        file_config = json.loads(config_path.read_text(encoding="utf-8"))

    filters = file_config.get("filters", {})
    if isinstance(filters, list):
        filters = _parse_filter_args([str(item) for item in filters])
    elif not isinstance(filters, dict):
        filters = {}
    cli_filters = _parse_filter_args(args.filter)
    filters.update(cli_filters)

    criteria = args.criteria or file_config.get("criteria") or ""
    if not str(criteria).strip():
        raise ValueError("action screen-recommend-detail requires --criteria or config.criteria")

    default_report_path = str(Path.cwd() / "reports" / "recommend_detail_greeted.xlsx")
    report_path = args.report_path
    if args.config and args.report_path == default_report_path:
        report_path = file_config.get("reportPath") or file_config.get("report_path") or args.report_path

    return {
        "endpoint": args.endpoint or file_config.get("endpoint") or "http://127.0.0.1:9222",
        "mock_dir": args.mock_dir or file_config.get("mockDir") or file_config.get("mock_dir"),
        "url_contains": args.url_contains or file_config.get("urlContains") or file_config.get("url_contains") or "/web/chat/recommend",
        "job_title": args.job_title or file_config.get("jobTitle") or file_config.get("job_title") or "",
        "city": args.city or file_config.get("city") or "",
        "filters": filters,
        "criteria": str(criteria),
        "model": args.model or file_config.get("model"),
        "max_greetings": int(args.max_greetings or file_config.get("maxGreetings") or file_config.get("max_greetings") or 30),
        "max_checks": int(
            args.target_resumes
            or file_config.get("targetResumes")
            or file_config.get("target_resumes")
            or file_config.get("maxChecks")
            or file_config.get("max_checks")
            or args.max_checks
            or 200
        ),
        "post_filter_wait_seconds": float(
            file_config.get("postFilterWaitSeconds")
            or file_config.get("post_filter_wait_seconds")
            or args.post_filter_wait
            or 0
        ),
        "operation_delay_seconds": float(
            getattr(args, "operation_delay_seconds", None)
            if getattr(args, "operation_delay_seconds", None) is not None
            else (
                file_config.get("operationDelaySeconds")
                or file_config.get("operation_delay_seconds")
                or 0
            )
        ),
        "max_filter_retries": int(
            getattr(args, "max_filter_retries", None)
            if getattr(args, "max_filter_retries", None) is not None
            else (
                file_config.get("maxFilterRetries")
                or file_config.get("max_filter_retries")
                or 2
            )
        ),
        "report_path": report_path or default_report_path,
        "dry_run": bool(args.dry_run or file_config.get("dryRun") or file_config.get("dry_run")),
    }


def _screen_recommend_detail(
    client: ChromeDebugClient,
    url_contains: str,
    job_title: str,
    city: str,
    filters: dict[str, str],
    criteria: str,
    model: str | None,
    max_greetings: int,
    max_checks: int,
    post_filter_wait_seconds: float,
    report_path: Path,
    dry_run: bool,
    operation_delay_seconds: float = 0,
    max_filter_retries: int = 2,
) -> dict[str, Any]:
    def operation_delay() -> None:
        if operation_delay_seconds > 0:
            time.sleep(operation_delay_seconds)

    open_result = client.open_recommend_page(url_contains=url_contains)
    operation_delay()
    try:
        pre_filter_close = client.close_recommend_detail(url_contains=url_contains)
        operation_delay()
    except Exception as exc:
        pre_filter_close = {"closed": False, "reason": str(exc)}
    switch_job, search_result = _apply_recommend_job_and_filters(
        client=client,
        url_contains=url_contains,
        job_title=job_title,
        city=city,
        filters=filters,
        operation_delay=operation_delay,
        operation_delay_seconds=operation_delay_seconds,
        max_filter_retries=max_filter_retries,
    )
    operation_delay()
    hard_job_failure = bool(job_title.strip() and switch_job and switch_job.get("switched") is False)
    hard_filter_failure = bool(filters and search_result.get("filterVerificationPassed") is False)
    if hard_job_failure or hard_filter_failure:
        report = _write_recommend_detail_report(
            report_path=report_path,
            criteria=criteria,
            job_title=job_title,
            processed=[],
        )
        return {
            "criteria": criteria,
            "jobTitle": job_title,
            "filters": filters,
            "dryRun": dry_run,
            "maxGreetings": max_greetings,
            "targetResumes": max_checks,
            "postFilterWaitSeconds": post_filter_wait_seconds,
            "operationDelaySeconds": operation_delay_seconds,
            "checkedCount": 0,
            "matchedCount": 0,
            "greetedCount": 0,
            "matchedNames": [],
            "greetedNames": [],
            "openRecommendPage": open_result,
            "preFilterCloseDetail": pre_filter_close,
            "switchJob": switch_job,
            "search": search_result,
            "stable": None,
            "report": report,
            "processed": [],
            "source": {
                "url": "",
                "candidateCount": 0,
            },
            "reason": "job_filter_verification_failed" if hard_job_failure else "filter_verification_failed",
        }
    if post_filter_wait_seconds > 0:
        time.sleep(post_filter_wait_seconds)
    stable_result = client.wait_for_recommend_results_stable(url_contains=url_contains)
    candidates_payload = client.capture_recommend_candidates(
        max_count=max(max_checks * 10, max_checks),
        url_contains=url_contains,
    )

    processed: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    greeted_count = 0

    for item in candidates_payload.get("candidates", []):
        if len(processed) >= max_checks:
            break
        if filters and not _recommend_detail_matches_local_filters(item, {}, filters):
            continue
        candidate_key = _recommend_candidate_dedupe_key(item)
        if not candidate_key or candidate_key in seen_keys:
            continue
        seen_keys.add(candidate_key)

        # 兜底：进入下一个候选人前，先关闭可能残留的"为你推荐"弹窗
        pre_candidate_skip_result: dict[str, Any] | None = None
        try:
            pre_candidate_skip_result = client.dismiss_recommend_popup_cards(url_contains=url_contains)
            operation_delay()
        except Exception as exc:
            pre_candidate_skip_result = {"skipped": False, "dismissed": False, "reason": str(exc)}

        pre_close_result: dict[str, Any] | None = None
        try:
            pre_close_result = client.close_recommend_detail(url_contains=url_contains)
            operation_delay()
        except Exception as exc:
            pre_close_result = {"closed": False, "reason": str(exc)}

        if ((pre_close_result or {}).get("after") or {}).get("detailOpen"):
            open_candidate = {"opened": False, "reason": "previous_detail_not_closed"}
        else:
            open_candidate = client.open_recommend_candidate(
                candidate_id=str(item.get("candidateId") or ""),
                card_index=int(item.get("cardIndex") or 0),
                url_contains=url_contains,
            )
            operation_delay()
        resume: dict[str, Any] = {}
        analysis: dict[str, Any] = {}
        greet_result: dict[str, Any] | None = None
        close_result: dict[str, Any] | None = None
        post_close_skip_result: dict[str, Any] | None = None
        reason = ""

        try:
            if not open_candidate.get("opened"):
                reason = str(open_candidate.get("reason") or "candidate_not_opened")
            else:
                operation_delay()
                resume = client.capture_candidate_resume(url_contains=url_contains)
                resume_text = str(resume.get("rawTextPreview") or "")
                resume_preview_text = resume_text[:6000]
                enriched_key = _recommend_candidate_dedupe_key(item, resume=resume)
                if enriched_key and enriched_key != candidate_key:
                    if enriched_key in seen_keys:
                        reason = "duplicate_candidate"
                        resume_text = ""
                        resume_preview_text = ""
                    else:
                        seen_keys.add(enriched_key)
                previous_text = str(
                    ((pre_close_result or {}).get("before") or {}).get("textPreview")
                    or ((pre_close_result or {}).get("after") or {}).get("textPreview")
                    or ""
                )
                if reason == "duplicate_candidate":
                    pass
                elif previous_text.strip():
                    previous_signature = re.sub(r"\s+", "", previous_text)[:120]
                    stale_deadline = time.time() + 6.0
                    while previous_signature and re.sub(r"\s+", "", resume_text).startswith(previous_signature) and time.time() < stale_deadline:
                        time.sleep(0.5)
                        resume = client.capture_candidate_resume(url_contains=url_contains)
                        resume_text = str(resume.get("rawTextPreview") or "")
                        resume_preview_text = resume_text[:6000]
                if reason == "duplicate_candidate":
                    pass
                elif not resume.get("detailOpen"):
                    reason = "resume_detail_not_open"
                elif resume.get("captureStatus") in {"detail_still_loading", "canvas_not_ready"}:
                    reason = str(resume.get("captureStatus"))
                elif not resume_text.strip():
                    reason = "resume_not_captured"
                elif previous_text.strip() and re.sub(r"\s+", "", resume_text).startswith(re.sub(r"\s+", "", previous_text)[:120]):
                    reason = "resume_detail_stale"
                elif not _recommend_detail_matches_local_filters(item, resume, filters):
                    reason = "page_filter_not_met"
                else:
                    analysis = analyze_resume_against_criteria(
                        resume_text=resume_preview_text,
                        criteria=criteria,
                        model=model,
                    )
                    analysis = {**analysis, "analysisInputSource": "resume_preview", "analysisInputLength": len(resume_preview_text)}
                    confidence_level = str(analysis.get("confidenceLevel") or "").lower()
                    llm_fallback_error = str(analysis.get("agent") or "") == "rules_fallback_after_llm_error"
                    high_confidence = not confidence_level or confidence_level == "high"
                    if llm_fallback_error:
                        reason = "llm_error_no_auto_greet"
                    elif analysis.get("meetsCriteria") and high_confidence:
                        if dry_run:
                            greet_result = {"greeted": False, "dryRun": True}
                        else:
                            greet_result = _click_recommend_card_greet_with_retries(
                                client=client,
                                item=item,
                                url_contains=url_contains,
                                attempts=3,
                                operation_delay_seconds=operation_delay_seconds,
                            )
                            operation_delay()
                            if greet_result.get("greeted"):
                                greeted_count += 1
                    elif analysis.get("meetsCriteria"):
                        reason = "criteria_low_confidence"
                    else:
                        reason = "criteria_not_met"
        except Exception as exc:
            reason = "error"
            analysis = {**analysis, "error": str(exc)}
        finally:
            try:
                close_result = client.close_recommend_detail(url_contains=url_contains)
                operation_delay()
            except Exception as exc:
                close_result = {"closed": False, "reason": str(exc)}
            # 打招呼后关闭"为你推荐"弹窗，避免遮挡下一个候选人
            try:
                post_close_skip_result = client.dismiss_recommend_popup_cards(url_contains=url_contains)
                operation_delay()
            except Exception as exc:
                post_close_skip_result = {"skipped": False, "dismissed": False, "reason": str(exc)}

        candidate_name = (
            analysis.get("name")
            or resume.get("name")
            or resume.get("candidateName")
            or item.get("candidateName")
            or ""
        )
        processed.append(
            {
                "checkedAt": datetime.now().isoformat(timespec="seconds"),
                "candidate": item,
                "preCandidateSkipSimilarRecommend": pre_candidate_skip_result,
                "preCloseDetail": pre_close_result,
                "openCandidate": open_candidate,
                "resume": resume,
                "analysis": analysis,
                "greet": greet_result,
                "closeDetail": close_result,
                "postCloseSkipSimilarRecommend": post_close_skip_result,
                "name": candidate_name,
                "meetsCriteria": bool(analysis.get("meetsCriteria")),
                "confidenceScore": analysis.get("confidenceScore", ""),
                "confidenceLevel": analysis.get("confidenceLevel", ""),
                "confidenceReason": analysis.get("confidenceReason", ""),
                "greeted": bool(greet_result and greet_result.get("greeted")),
                "reason": reason,
            }
        )

        if greeted_count >= max_greetings:
            break

    report = _write_recommend_detail_report(
        report_path=report_path,
        criteria=criteria,
        job_title=job_title,
        processed=processed,
    )
    matched = [item for item in processed if item.get("meetsCriteria")]
    greeted = [item for item in processed if item.get("greeted")]
    return {
        "criteria": criteria,
        "jobTitle": job_title,
        "filters": filters,
        "dryRun": dry_run,
        "maxGreetings": max_greetings,
        "targetResumes": max_checks,
        "postFilterWaitSeconds": post_filter_wait_seconds,
        "operationDelaySeconds": operation_delay_seconds,
        "checkedCount": len(processed),
        "matchedCount": len(matched),
        "greetedCount": len(greeted),
        "matchedNames": [str(item.get("name") or "") for item in matched if item.get("name")],
        "greetedNames": [str(item.get("name") or "") for item in greeted if item.get("name")],
        "openRecommendPage": open_result,
        "preFilterCloseDetail": pre_filter_close,
        "switchJob": switch_job,
        "search": search_result,
        "stable": stable_result,
        "report": report,
        "processed": processed,
        "source": {
            "url": candidates_payload.get("url", ""),
            "candidateCount": candidates_payload.get("count", 0),
        },
    }


def _apply_recommend_job_and_filters(
    client: ChromeDebugClient,
    url_contains: str,
    job_title: str,
    city: str,
    filters: dict[str, str],
    operation_delay: Any,
    operation_delay_seconds: float,
    max_filter_retries: int,
) -> tuple[dict[str, Any] | None, dict[str, Any]]:
    attempts: list[dict[str, Any]] = []
    switch_job: dict[str, Any] | None = None
    search_result: dict[str, Any] = {}
    max_attempts = max(1, max_filter_retries + 1)

    for attempt in range(1, max_attempts + 1):
        if job_title.strip():
            switch_job = client.switch_recommend_job_filter(job_text=job_title, url_contains=url_contains)
            operation_delay()

        search_result = client.apply_recommend_filters(
            job_title="",
            city=city,
            filters=filters,
            url_contains=url_contains,
            operation_delay_seconds=operation_delay_seconds,
            max_filter_retries=max_filter_retries,
        )
        operation_delay()

        job_failed = bool(job_title.strip() and switch_job and switch_job.get("switched") is False)
        attempts.append(
            {
                "attempt": attempt,
                "switchJob": switch_job,
                "search": search_result,
                "reason": "job_filter_not_switched" if job_failed else "",
            }
        )
        if not job_failed:
            break

    if switch_job is not None:
        switch_job = {**switch_job, "retryHistory": attempts[:-1], "attemptCount": len(attempts)}
    search_result = {
        **search_result,
        "jobFilterRetryHistory": attempts[:-1],
        "jobFilterAttemptCount": len(attempts),
    }
    return switch_job, search_result


def _recommend_candidate_dedupe_key(
    candidate: dict[str, Any],
    resume: dict[str, Any] | None = None,
    analysis: dict[str, Any] | None = None,
) -> str:
    resume = resume or {}
    analysis = analysis or {}
    name = (
        analysis.get("name")
        or resume.get("name")
        or resume.get("candidateName")
        or candidate.get("candidateName")
        or ""
    )
    salary = candidate.get("salary") or ""
    profile = (
        candidate.get("experience")
        or candidate.get("age")
        or resume.get("yearsOfExperience")
        or ""
    )
    expectation = (
        resume.get("expectedTitle")
        or candidate.get("expectedCity")
        or candidate.get("expectedTitle")
        or ""
    )
    preview = str(candidate.get("detailTextPreview") or resume.get("rawTextPreview") or "")[:180]
    parts = [name, salary, profile, expectation, preview]
    normalized = [
        re.sub(r"\s+", "", str(part or "")).lower()
        for part in parts
        if str(part or "").strip()
    ]
    if normalized:
        return "::".join(normalized)
    candidate_id = str(candidate.get("candidateId") or "")
    if candidate_id and not re.fullmatch(r"(?:candidate-)?\d+", candidate_id):
        return f"id::{candidate_id}"
    return ""


def _click_recommend_card_greet_with_retries(
    client: ChromeDebugClient,
    item: dict[str, Any],
    url_contains: str,
    attempts: int = 4,
    operation_delay_seconds: float = 0,
) -> dict[str, Any]:
    def operation_delay() -> None:
        if operation_delay_seconds > 0:
            time.sleep(operation_delay_seconds)

    results: list[dict[str, Any]] = []
    final_result: dict[str, Any] = {"greeted": False, "clicked": False, "attempts": results}
    for attempt in range(1, max(1, attempts) + 1):
        # Dismiss any popup cards before attempting
        try:
            pre_popup_dismiss = client.dismiss_recommend_popup_cards(url_contains=url_contains)
        except Exception as exc:
            pre_popup_dismiss = {"dismissed": False, "reason": str(exc)}
        try:
            close_before_greet = client.close_recommend_detail(url_contains=url_contains)
            operation_delay()
        except Exception as exc:
            close_before_greet = {"closed": False, "reason": str(exc)}
        try:
            result = client.click_recommend_card_greet(
                candidate_id=str(item.get("candidateId") or ""),
                card_index=int(item.get("cardIndex") or 0),
                url_contains=url_contains,
            )
            operation_delay()
        except Exception as exc:
            result = {"greeted": False, "clicked": False, "reason": str(exc)}
        result = {
            **result,
            "attempt": attempt,
            "method": result.get("method") or "card_greet_after_detail_close",
            "closeBeforeGreet": close_before_greet,
            "prePopupDismiss": pre_popup_dismiss,
        }
        if not result.get("greeted"):
            try:
                pre_detail_fallback_skip = client.dismiss_recommend_popup_cards(url_contains=url_contains)
                time.sleep(0.3)
                detail_fallback = client.click_recommend_detail_greet(url_contains=url_contains)
                operation_delay()
            except Exception as exc:
                pre_detail_fallback_skip = {"skipped": False, "dismissed": False, "reason": str(exc)}
                detail_fallback = {"greeted": False, "clicked": False, "reason": str(exc)}
            result["preDetailFallbackSkipSimilarRecommend"] = pre_detail_fallback_skip
            result["detailFallback"] = detail_fallback
            if detail_fallback.get("greeted"):
                result["greeted"] = True
        results.append(result)
        final_result = {**result, "attempts": results, "retryCount": attempt - 1}
        if result.get("greeted"):
            try:
                post_greet_skip = client.dismiss_recommend_popup_cards(url_contains=url_contains)
            except Exception as exc:
                post_greet_skip = {"skipped": False, "dismissed": False, "reason": str(exc)}
            final_result["postGreetSkipSimilarRecommend"] = post_greet_skip
            result["postGreetSkipSimilarRecommend"] = post_greet_skip
            return final_result
        time.sleep(1.5)
    return final_result


def _recommend_detail_matches_local_filters(
    candidate: dict[str, Any],
    resume: dict[str, Any],
    filters: dict[str, Any],
) -> bool:
    combined = " ".join(
        str(value)
        for value in [
            candidate.get("salary"),
            candidate.get("age"),
            candidate.get("education"),
            candidate.get("detailTextPreview"),
            resume.get("education"),
            resume.get("yearsOfExperience"),
            resume.get("rawTextPreview"),
        ]
        if value
    )
    return _recommend_detail_matches_local_filters_v2(combined, filters)


def _recommend_detail_matches_local_filters_v2(combined: str, filters: dict[str, Any]) -> bool:
    for key, value in filters.items():
        normalized_key = str(key)
        normalized_values = _split_filter_values(value)
        if any(label in normalized_key for label in ("学历", "学历要求", "教育")):
            if normalized_values and not any(item in combined for item in normalized_values):
                return False
        elif any(label in normalized_key for label in ("年龄", "年纪", "age")):
            age_ranges = [item for item in (_parse_numeric_range(raw) for raw in normalized_values) if item]
            age = _first_int_matching(combined, r"(\d{1,2})\s*岁")
            if age_ranges and (age is None or not any(lower <= age <= upper for lower, upper in age_ranges)):
                return False
        elif any(label in normalized_key for label in ("薪资", "薪资待遇", "待遇")):
            target_ranges = [item for item in (_parse_numeric_range(raw) for raw in normalized_values) if item]
            salary_range = _parse_salary_range(combined)
            if target_ranges and salary_range and not any(_ranges_overlap(target_range, salary_range) for target_range in target_ranges):
                return False
        elif any(label in normalized_key for label in ("\u7ecf\u9a8c", "\u5de5\u4f5c\u7ecf\u9a8c")):
            target_ranges = [item for item in (_parse_numeric_range(raw) for raw in normalized_values) if item]
            work_years = _first_int_matching(combined, r"(\d{1,2})\s*\u5e74")
            if target_ranges and work_years is not None and not any(lower <= work_years <= upper for lower, upper in target_ranges):
                return False
    return True


def _split_filter_values(value: Any) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        raw_items = [str(item) for item in value]
    else:
        raw_items = re.split(r"[,，、;；|/]|(?:\s+or\s+)|或", str(value or ""), flags=re.IGNORECASE)
    return [item.strip() for item in raw_items if item.strip()]


def _parse_numeric_range(value: str) -> tuple[int, int] | None:
    numbers = [int(item) for item in re.findall(r"\d+", str(value))]
    if len(numbers) < 2:
        return None
    lower, upper = numbers[0], numbers[1]
    if lower > upper:
        lower, upper = upper, lower
    return lower, upper


def _parse_salary_range(value: str) -> tuple[int, int] | None:
    match = re.search(r"(\d+)\s*-\s*(\d+)\s*[kK]", str(value))
    if not match:
        return None
    lower, upper = int(match.group(1)), int(match.group(2))
    if lower > upper:
        lower, upper = upper, lower
    return lower, upper


def _ranges_overlap(left: tuple[int, int] | dict[str, Any], right: tuple[int, int] | dict[str, Any]) -> bool:
    def bounds(value: tuple[int, int] | dict[str, Any]) -> tuple[float, float]:
        if isinstance(value, dict):
            return float(value.get("min", 0)), float(value.get("max", 0))
        return float(value[0]), float(value[1])

    left_min, left_max = bounds(left)
    right_min, right_max = bounds(right)
    return left_min <= right_max and right_min <= left_max


def _first_int_matching(value: str, pattern: str) -> int | None:
    match = re.search(pattern, value)
    return int(match.group(1)) if match else None


def _write_recommend_detail_report(
    report_path: Path,
    criteria: str,
    job_title: str,
    processed: list[dict[str, Any]],
) -> dict[str, Any]:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    sorted_items = sorted(
        processed,
        key=lambda item: (
            not bool(item.get("greeted")),
            not bool(item.get("meetsCriteria")),
            str(item.get("name") or ""),
        ),
    )
    rows = []
    for item in sorted_items:
        analysis = item.get("analysis") or {}
        resume = item.get("resume") or {}
        greet = item.get("greet") or {}
        rows.append(
            {
                "checked_at": item.get("checkedAt", ""),
                "job_title": job_title,
                "candidate": item.get("name", ""),
                "meets_criteria": "yes" if item.get("meetsCriteria") else "no",
                "confidence_score": analysis.get("confidenceScore", item.get("confidenceScore", "")),
                "confidence_level": analysis.get("confidenceLevel", item.get("confidenceLevel", "")),
                "confidence_reason": analysis.get("confidenceReason", item.get("confidenceReason", "")),
                "experience_evidence": _format_experience_evidence(analysis.get("experienceEvidence") or []),
                "rejected_evidence": _format_rejected_evidence(analysis.get("rejectedEvidence") or []),
                "greeted": "yes" if item.get("greeted") else "no",
                "greet_status": json.dumps(greet, ensure_ascii=False) if greet else item.get("reason", ""),
                "candidate_summary": analysis.get("candidateSummary", ""),
                "reasons": "；".join(analysis.get("reasons", [])),
                "risks": "；".join(analysis.get("risks", [])),
                "keyword_evidence": _format_keyword_evidence(analysis),
                "criteria": criteria,
                "agent": analysis.get("agent", ""),
                "analysis_input_source": analysis.get("analysisInputSource", ""),
                "analysis_input_length": analysis.get("analysisInputLength", ""),
                "resume_source": resume.get("resumeSource") or ("detail" if resume.get("detailOpen") else "not_captured"),
                "ocr_status": resume.get("ocrStatus", ""),
                "ocr_error": str((resume.get("ocr") or {}).get("error") or ""),
                "resume_preview": str(resume.get("rawTextPreview") or "")[:6000],
            }
        )

    if report_path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "recommend_detail"
            headers = list(rows[0].keys()) if rows else [
                "checked_at", "job_title", "candidate", "meets_criteria", "greeted",
                "confidence_score", "confidence_level", "confidence_reason", "experience_evidence", "rejected_evidence",
                "greet_status", "candidate_summary", "reasons", "risks", "keyword_evidence", "criteria",
                "agent", "analysis_input_source", "analysis_input_length",
                "resume_source", "ocr_status", "ocr_error", "resume_preview"
            ]
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append([row.get(header, "") for header in headers])
            for idx, header in enumerate(headers, start=1):
                max_value = max([len(str(header))] + [len(str(row.get(header, ""))) for row in rows[:20]])
                sheet.column_dimensions[get_column_letter(idx)].width = min(60, max(12, max_value + 2))
            workbook.save(report_path)
            return {"path": str(report_path), "format": "xlsx", "rowCount": len(rows)}
        except Exception as exc:
            csv_path = report_path.with_suffix(".csv")
            _write_rows_csv(csv_path, rows)
            return {"path": str(csv_path), "format": "csv", "rowCount": len(rows), "xlsxError": str(exc)}

    _write_rows_csv(report_path, rows)
    return {"path": str(report_path), "format": "csv", "rowCount": len(rows)}


def _format_keyword_evidence(analysis: dict[str, Any]) -> str:
    matches = ((analysis.get("keywordMatch") or {}).get("matches") or [])
    pieces: list[str] = []
    for match in matches:
        keyword = str(match.get("keyword") or "").strip()
        matched_term = str(match.get("matchedTerm") or "").strip()
        summary = str(match.get("evidenceSummary") or "").strip()
        evidence = str(match.get("evidenceText") or "").strip()
        if not (keyword or matched_term or summary or evidence):
            continue
        label = f"{keyword} -> {matched_term}".strip(" ->")
        detail = summary or evidence
        if evidence and evidence not in detail:
            detail = f"{detail}（{evidence}）" if detail else evidence
        pieces.append(f"{label}：{detail}" if label and detail else label or detail)
    return "；".join(pieces)

def _format_experience_evidence(items: list[dict[str, Any]]) -> str:
    pieces: list[str] = []
    for item in items:
        company = str(item.get("company") or "").strip()
        title = str(item.get("title") or "").strip()
        start = str(item.get("start") or "").strip()
        end = str(item.get("end") or "").strip()
        months = item.get("months", "")
        terms = "、".join(str(term) for term in (item.get("platformTerms") or item.get("roleTerms") or [])[:5])
        text = str(item.get("evidenceText") or "").strip()
        head = " ".join(part for part in [company, title, f"{start}-{end}".strip("-")] if part)
        detail = f"{months}个月" if months != "" else ""
        if terms:
            detail = f"{detail} 命中:{terms}".strip()
        if text:
            detail = f"{detail} {text[:120]}".strip()
        pieces.append(f"{head} | {detail}".strip(" |"))
    return "；".join(pieces)


def _format_rejected_evidence(items: list[dict[str, Any]]) -> str:
    pieces: list[str] = []
    for item in items:
        reason = str(item.get("reason") or "").strip()
        raw = str(item.get("raw") or item.get("text") or item.get("evidenceText") or "").strip()
        if reason or raw:
            pieces.append(f"{reason}: {raw[:120]}".strip(": "))
    return "；".join(pieces)


def _screen_and_greet_recommendations(
    client: ChromeDebugClient,
    url_contains: str,
    criteria: str,
    max_scrolls: int,
    model: str | None,
    max_greetings: int,
    dry_run: bool,
) -> dict[str, Any]:
    cards_payload = client.capture_recommend_resume_cards(
        max_scrolls=max_scrolls,
        url_contains=url_contains,
    )
    analyzed: list[dict[str, Any]] = []
    seen_names: set[str] = set()

    for card in cards_payload.get("cards", []):
        analysis = analyze_resume_against_criteria(
            resume_text=str(card.get("rawText") or ""),
            criteria=criteria,
            model=model,
        )
        name = str(analysis.get("name") or card.get("name") or "").strip()
        row = {
            "name": name,
            "cardSummary": card.get("summary", ""),
            "candidateSummary": analysis.get("candidateSummary", ""),
            "meetsCriteria": analysis.get("meetsCriteria", False),
            "educationLevel": analysis.get("educationLevel", ""),
            "workYears": analysis.get("workYears"),
            "agent": analysis.get("agent", ""),
            "reasons": analysis.get("reasons", []),
            "risks": analysis.get("risks", []),
        }
        analyzed.append(row)

    matched: list[dict[str, Any]] = []
    for item in analyzed:
        name = str(item.get("name", "")).strip()
        if not name or name in seen_names or not item.get("meetsCriteria"):
            continue
        seen_names.add(name)
        matched.append(item)

    greeting_names = [str(item["name"]) for item in matched[:max_greetings]]
    greet_result = {"results": [], "dryRun": True}
    if greeting_names and not dry_run:
        greet_result = client.greet_recommend_candidates_by_name(
            names=greeting_names,
            url_contains=url_contains,
        )

    return {
        "criteria": criteria,
        "checkedCount": len(analyzed),
        "matchedCount": len(matched),
        "matchedNames": [item["name"] for item in matched],
        "greetingNames": greeting_names,
        "dryRun": dry_run,
        "greetResult": greet_result,
        "matched": matched,
        "analyzed": analyzed,
        "source": {
            "url": cards_payload.get("url", ""),
            "frameUrl": cards_payload.get("frameUrl", ""),
            "cardCount": cards_payload.get("count", 0),
        },
    }


def _default_capability_report_path() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path.cwd() / "reports" / f"recommend_capability_{timestamp}.xlsx"


def _screen_and_greet_capability(
    client: ChromeDebugClient,
    url_contains: str,
    criteria: str,
    max_scrolls: int,
    max_greetings: int,
    report_path: Path,
    dry_run: bool,
) -> dict[str, Any]:
    cards_payload = client.capture_recommend_resume_cards(
        max_scrolls=max_scrolls,
        url_contains=url_contains,
    )
    analyzed: list[dict[str, Any]] = []
    seen_names: set[str] = set()
    matched_names: list[str] = []

    for card in cards_payload.get("cards", []):
        name = str(card.get("name") or "").strip()
        analysis = analyze_operations_capability(
            resume_text=str(card.get("rawText") or ""),
            candidate_name=name,
            criteria=criteria,
        )
        row = {
            "name": analysis.get("name") or name,
            "summary": card.get("summary", ""),
            "candidateSummary": analysis.get("candidateSummary", ""),
            "meetsCriteria": analysis.get("meetsCriteria", False),
            "capabilityCount": analysis.get("capabilityCount", 0),
            "matchedCapabilities": analysis.get("matchedCapabilities", []),
            "projectEvidence": analysis.get("projectEvidence", []),
            "reasons": analysis.get("reasons", []),
            "risks": analysis.get("risks", []),
            "agent": analysis.get("agent", ""),
            "rawTextPreview": str(card.get("rawText") or "")[:600],
        }
        analyzed.append(row)
        row_name = str(row.get("name") or "").strip()
        if row.get("meetsCriteria") and row_name and row_name not in seen_names:
            seen_names.add(row_name)
            matched_names.append(row_name)

    greeting_names = matched_names[:max_greetings]
    greet_result = {"results": [], "dryRun": True}
    if greeting_names and not dry_run:
        greet_result = client.greet_recommend_candidates_by_name(
            names=greeting_names,
            url_contains=url_contains,
        )

    report = _write_capability_report(
        report_path=report_path,
        criteria=criteria,
        analyzed=analyzed,
        greet_result=greet_result,
        dry_run=dry_run,
    )
    return {
        "criteria": criteria,
        "checkedCount": len(analyzed),
        "matchedCount": len(matched_names),
        "matchedNames": matched_names,
        "greetingNames": greeting_names,
        "dryRun": dry_run,
        "greetResult": greet_result,
        "report": report,
        "source": {
            "url": cards_payload.get("url", ""),
            "frameUrl": cards_payload.get("frameUrl", ""),
            "cardCount": cards_payload.get("count", 0),
        },
    }


def _write_capability_report(
    report_path: Path,
    criteria: str,
    analyzed: list[dict[str, Any]],
    greet_result: dict[str, Any],
    dry_run: bool,
) -> dict[str, Any]:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    greeted_by_name = {
        str(item.get("name") or ""): item
        for item in greet_result.get("results", [])
        if item.get("greeted")
    }
    report_items = [
        item for item in analyzed
        if item.get("meetsCriteria") and str(item.get("name") or "") in greeted_by_name
    ]
    if dry_run:
        report_items = [item for item in analyzed if item.get("meetsCriteria")]
    rows = []
    for item in report_items:
        greet = greeted_by_name.get(str(item.get("name") or ""), {})
        rows.append(
            {
                "候选人": item.get("name", ""),
                "是否满足": "是" if item.get("meetsCriteria") else "否",
                "能力匹配数": item.get("capabilityCount", 0),
                "匹配能力": "、".join(item.get("matchedCapabilities", [])),
                "选择原因": "；".join(item.get("reasons", [])),
                "风险/缺口": "；".join(item.get("risks", [])),
                "项目证据": "；".join(item.get("projectEvidence", [])),
                "摘要": item.get("candidateSummary", ""),
                "打招呼状态": "dry_run" if dry_run else json.dumps(greet, ensure_ascii=False),
                "筛选条件": criteria,
                "简历片段": item.get("rawTextPreview", ""),
            }
        )

    if report_path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "筛选结果"
            headers = list(rows[0].keys()) if rows else [
                "候选人", "是否满足", "能力匹配数", "匹配能力", "选择原因", "风险/缺口", "项目证据", "摘要", "打招呼状态", "筛选条件", "简历片段"
            ]
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append([row.get(header, "") for header in headers])
            for idx, header in enumerate(headers, start=1):
                width = min(60, max(12, len(header) + 4))
                sheet.column_dimensions[get_column_letter(idx)].width = width
            workbook.save(report_path)
            return {"path": str(report_path), "format": "xlsx", "rowCount": len(rows)}
        except Exception as exc:
            csv_path = report_path.with_suffix(".csv")
            _write_rows_csv(csv_path, rows)
            return {"path": str(csv_path), "format": "csv", "rowCount": len(rows), "xlsxError": str(exc)}

    _write_rows_csv(report_path, rows)
    return {"path": str(report_path), "format": "csv", "rowCount": len(rows)}


def _write_rows_csv(report_path: Path, rows: list[dict[str, Any]]) -> None:
    headers = list(rows[0].keys()) if rows else ["候选人", "是否满足", "选择原因"]
    with report_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def _request_resume_new_greetings(
    client: ChromeDebugClient,
    url_contains: str,
    job_text: str,
    message: str,
    max_count: int,
    inbox_scrolls: int,
    wait_after_unread: float,
    operation_delay_seconds: float,
    send: bool,
    criteria: str = "",
) -> dict[str, Any]:
    dry_run = not send
    if send and not criteria.strip():
        return _request_resume_new_greetings_stopped(
            job_text, message, criteria, dry_run, inbox_scrolls, wait_after_unread, operation_delay_seconds,
            None, None, None, None, "criteria_required",
        )

    switch_status = client.switch_conversation_status_tab(label="\u65b0\u62db\u547c", url_contains=url_contains)
    if not switch_status.get("switched"):
        return _request_resume_new_greetings_stopped(
            job_text, message, criteria, dry_run, inbox_scrolls, wait_after_unread, operation_delay_seconds,
            switch_status, None, None, None, "status_tab_not_switched",
        )
    _sleep_operation_delay(operation_delay_seconds)

    switch_job = client.switch_job_filter(job_text=job_text, url_contains=url_contains)
    if not switch_job.get("switched"):
        return _request_resume_new_greetings_stopped(
            job_text, message, criteria, dry_run, inbox_scrolls, wait_after_unread, operation_delay_seconds,
            switch_status, switch_job, None, None, "job_filter_not_switched",
        )
    _sleep_operation_delay(operation_delay_seconds)

    switch_unread = client.switch_message_filter(label="\u672a\u8bfb", url_contains=url_contains)
    _sleep_operation_delay(operation_delay_seconds)
    if wait_after_unread > 0:
        time.sleep(wait_after_unread)
    inbox = client.capture_unread_inbox_scrolled(
        url_contains=url_contains,
        max_items=max_count,
        max_scrolls=inbox_scrolls,
    )
    conversations = [item for item in inbox.get("conversations", []) if item.get("dataId")][:max_count]
    results: list[dict[str, Any]] = []

    for item in conversations:
        if dry_run:
            results.append(
                {
                    "target": item,
                    "openConversation": None,
                    "conversationReady": None,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": None,
                    "onlineResumeOpen": None,
                    "resume": None,
                    "analysis": None,
                    "closeAttachmentPreview": None,
                    "success": False,
                    "skipped": True,
                    "reason": "dry_run",
                }
            )
            _sleep_operation_delay(operation_delay_seconds)
            continue

        open_result = client.open_conversation(
            data_id=item.get("dataId"),
            candidate_name=item.get("candidateName"),
            url_contains=url_contains,
            max_scrolls=inbox_scrolls,
        )
        _sleep_operation_delay(operation_delay_seconds)
        readiness = client.wait_for_conversation_ready(
            expected_candidate_name=item.get("candidateName"),
            expected_job_title=item.get("jobTitle"),
            url_contains=url_contains,
        )
        if not readiness.get("ready"):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": None,
                    "onlineResumeOpen": None,
                    "resume": None,
                    "analysis": None,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": "conversation_not_ready",
                }
            )
            continue

        conversation = readiness.get("conversation", {})
        candidate_precheck = _analyze_candidate_precheck(
            conversation,
            job_text=job_text,
        )
        if candidate_precheck.get("skip"):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": candidate_precheck,
                    "onlineResumeOpen": None,
                    "resume": None,
                    "analysis": None,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": "candidate_precheck_skipped",
                }
            )
            continue

        if _conversation_has_resume_request_or_attachment(conversation):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": candidate_precheck,
                    "onlineResumeOpen": None,
                    "resume": None,
                    "analysis": None,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": "resume_already_requested_or_present",
                }
            )
            continue

        online_resume_open = client.open_online_resume(url_contains=url_contains)
        _sleep_operation_delay(operation_delay_seconds)
        if not online_resume_open.get("opened"):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": candidate_precheck,
                    "onlineResumeOpen": online_resume_open,
                    "resume": None,
                    "analysis": None,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": "online_resume_not_found",
                }
            )
            continue

        resume = client.capture_candidate_resume(url_contains=url_contains)
        _sleep_operation_delay(operation_delay_seconds)
        resume_text = str(resume.get("rawTextPreview") or "").strip()
        if not resume_text:
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": candidate_precheck,
                    "onlineResumeOpen": online_resume_open,
                    "resume": resume,
                    "analysis": None,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": "resume_not_captured",
                }
            )
            continue

        analysis = analyze_resume_against_criteria(
            resume_text=resume_text[:6000],
            criteria=criteria,
            model=None,
        )
        analysis = {**analysis, "analysisInputSource": "online_resume_ocr", "analysisInputLength": len(resume_text[:6000])}
        confidence_level = str(analysis.get("confidenceLevel") or "").lower()
        llm_fallback_error = str(analysis.get("agent") or "") == "rules_fallback_after_llm_error"
        if llm_fallback_error or not analysis.get("meetsCriteria") or (confidence_level and confidence_level != "high"):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "draftFill": None,
                    "send": None,
                    "requestResume": None,
                    "candidatePrecheck": candidate_precheck,
                    "onlineResumeOpen": online_resume_open,
                    "resume": resume,
                    "analysis": analysis,
                    "closeAttachmentPreview": close_result,
                    "success": False,
                    "skipped": True,
                    "reason": (
                        "llm_error_no_auto_send"
                        if llm_fallback_error
                        else "criteria_low_confidence" if analysis.get("meetsCriteria") else "criteria_not_met"
                    ),
                }
            )
            continue

        close_result = client.close_attachment_preview(url_contains=url_contains)
        _sleep_operation_delay(operation_delay_seconds)
        common_phrase = client.click_first_common_phrase(url_contains=url_contains)
        _sleep_operation_delay(operation_delay_seconds)
        reason = ""
        send_result = None
        if common_phrase.get("clicked"):
            send_result = client.send_current_message(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            if not send_result.get("sent"):
                reason = str(send_result.get("reason") or "send_failed")
        else:
            reason = str(common_phrase.get("reason") or "common_phrase_not_clicked")
        success = bool(common_phrase.get("clicked") and send_result and send_result.get("sent"))
        results.append(
            {
                "target": item,
                "openConversation": open_result,
                "conversationReady": readiness,
                "fill": None,
                "draftFill": None,
                "commonPhrase": common_phrase,
                "send": send_result,
                "requestResume": None,
                "candidatePrecheck": candidate_precheck,
                "onlineResumeOpen": online_resume_open,
                "resume": resume,
                "analysis": analysis,
                "closeAttachmentPreview": close_result,
                "success": success,
                "skipped": not success,
                "reason": "" if success else reason,
            }
        )

    return _request_resume_new_greetings_payload(
        job_text, message, criteria, dry_run, inbox_scrolls, wait_after_unread, operation_delay_seconds,
        switch_status, switch_job, switch_unread, inbox, conversations, results,
    )


def _sleep_operation_delay(operation_delay_seconds: float) -> None:
    if operation_delay_seconds > 0:
        time.sleep(operation_delay_seconds)


def _analyze_candidate_precheck(
    conversation: dict[str, Any],
    today: date | None = None,
    job_text: str = "",
) -> dict[str, Any]:
    today = today or date.today()
    raw_text = _candidate_profile_text(conversation)
    resolved_job_text = _conversation_job_text(conversation, job_text=job_text)
    fresh_graduate_exempt = _is_intern_or_assistant_job(resolved_job_text)
    reasons: list[str] = []
    evidence: list[dict[str, Any]] = []

    job_salary = _parse_salary_range(job_text)
    candidate_salary = _parse_salary_range(str(conversation.get("salary") or raw_text))
    if job_salary and candidate_salary and candidate_salary["max"] < job_salary["min"]:
        reasons.append("candidate_salary_below_job_range")
        evidence.append(
            {
                "type": "candidate_salary_below_job_range",
                "jobSalary": job_salary,
                "candidateSalary": candidate_salary,
            }
        )

    fresh_markers = ["应届", "在校", "25年毕业", "26年毕业"]
    matched_fresh = [marker for marker in fresh_markers if marker in raw_text]
    matched_fresh.extend(re.findall(r"\d{2}年(?:应届|毕业)", raw_text))
    matched_fresh = list(dict.fromkeys(matched_fresh))
    if matched_fresh and not fresh_graduate_exempt:
        reasons.append("fresh_graduate")
        evidence.append({"type": "fresh_graduate", "markers": matched_fresh})
    elif matched_fresh:
        evidence.append(
            {
                "type": "fresh_graduate_ignored_for_job",
                "markers": matched_fresh,
                "jobText": resolved_job_text,
            }
        )

    periods = _extract_work_periods(raw_text)
    sorted_periods = sorted(periods, key=lambda item: item["startIndex"])
    for previous, current in zip(sorted_periods, sorted_periods[1:]):
        previous_end = previous.get("endIndex")
        if previous_end is None:
            continue
        gap_months = int(current["startIndex"]) - int(previous_end) - 1
        if gap_months >= 6:
            reasons.append("work_gap_6_months_or_more")
            evidence.append(
                {
                    "type": "work_gap_between_jobs",
                    "gapMonths": gap_months,
                    "previous": previous["raw"],
                    "current": current["raw"],
                }
            )
            break

    if sorted_periods:
        latest = max(sorted_periods, key=lambda item: item.get("endIndex") if item.get("endIndex") is not None else 10**9)
        latest_end = latest.get("endIndex")
        if latest_end is not None:
            gap_months = _month_index(today.year, today.month) - int(latest_end)
            if gap_months >= 6:
                if "work_gap_6_months_or_more" not in reasons:
                    reasons.append("work_gap_6_months_or_more")
                evidence.append(
                    {
                        "type": "recent_work_gap",
                        "gapMonths": gap_months,
                        "latest": latest["raw"],
                        "today": today.isoformat(),
                    }
                )

    return {
        "skip": bool(reasons),
        "reasons": list(dict.fromkeys(reasons)),
        "evidence": evidence,
        "jobText": resolved_job_text,
        "jobSalary": job_salary,
        "candidateSalary": candidate_salary,
        "freshGraduateCheckExempt": fresh_graduate_exempt,
        "gapThresholdMonths": 6,
        "profileTimeline": {
            "rawText": raw_text,
            "workExperiences": [
                {
                    "raw": item["raw"],
                    "start": item["start"],
                    "end": item["end"],
                    "current": item["endIndex"] is None,
                }
                for item in sorted_periods
            ],
        },
    }


def _conversation_job_text(conversation: dict[str, Any], job_text: str = "") -> str:
    values = [
        str(job_text or ""),
        str(conversation.get("jobTitle") or ""),
        str(conversation.get("targetJobTitle") or ""),
    ]
    profile_text = _candidate_profile_text(conversation)
    match = re.search(r"沟通职位：\s*([^\n]+)", profile_text)
    if match:
        values.insert(0, match.group(1).strip())
    return " ".join(value.strip() for value in values if value and value.strip()).strip()


def _is_intern_or_assistant_job(job_text: str) -> bool:
    return any(marker in str(job_text or "") for marker in ("实习", "实习生", "助理"))


def _parse_salary_range(value: str) -> dict[str, float] | None:
    text = str(value or "")
    match = re.search(r"(?P<min>\d+(?:\.\d+)?)\s*-\s*(?P<max>\d+(?:\.\d+)?)\s*[kK]", text)
    if match:
        return {
            "raw": match.group(0),
            "min": float(match.group("min")),
            "max": float(match.group("max")),
            "unit": "K",
        }
    match = re.search(r"(?P<min>\d+(?:\.\d+)?)\s*-\s*(?P<max>\d+(?:\.\d+)?)\s*(?:千|k|K)", text)
    if match:
        return {
            "raw": match.group(0),
            "min": float(match.group("min")),
            "max": float(match.group("max")),
            "unit": "K",
        }
    return None


def _candidate_profile_text(conversation: dict[str, Any]) -> str:
    profile = conversation.get("profileTimeline") or {}
    values = [
        str(profile.get("rawText") or ""),
        str(conversation.get("candidateProfileText") or ""),
        str(conversation.get("sidebarText") or ""),
    ]
    unique_values = list(dict.fromkeys(value for value in values if value))
    return "\n".join(unique_values).strip()


def _extract_work_periods(raw_text: str) -> list[dict[str, Any]]:
    periods: list[dict[str, Any]] = []
    lines = [line.strip() for line in str(raw_text or "").splitlines() if line.strip()]
    period_pattern = re.compile(
        r"(?P<start>\d{4}(?:[./年]\d{1,2})?)\s*(?:-|－|–|—|~|至|到)\s*(?P<end>至今|现在|目前|今|Present|present|在职|\d{4}(?:[./年]\d{1,2})?)"
    )
    index = 0
    while index < len(lines):
        line = lines[index]
        matches = list(period_pattern.finditer(line))
        if _is_standalone_period_line(line, matches):
            date_lines: list[str] = []
            while index < len(lines):
                current_matches = list(period_pattern.finditer(lines[index]))
                if not _is_standalone_period_line(lines[index], current_matches):
                    break
                date_lines.append(lines[index])
                index += 1
            description_start = index
            for offset, date_line in enumerate(date_lines):
                context = lines[description_start + offset] if description_start + offset < len(lines) else date_line
                if _looks_like_education_line(context):
                    continue
                periods.extend(_periods_from_line(date_line, raw=context, period_pattern=period_pattern))
            continue

        value = line
        value = line.strip()
        if not value or _looks_like_education_line(value):
            index += 1
            continue
        periods.extend(_periods_from_line(value, raw=value, period_pattern=period_pattern))
        index += 1
    return periods


def _is_standalone_period_line(value: str, matches: list[re.Match[str]]) -> bool:
    if len(matches) != 1:
        return False
    remaining = value.replace(matches[0].group(0), "").strip()
    return not remaining


def _periods_from_line(
    value: str,
    raw: str,
    period_pattern: re.Pattern[str],
) -> list[dict[str, Any]]:
    periods: list[dict[str, Any]] = []
    for match in period_pattern.finditer(value):
        start = _parse_year_month(match.group("start"), default_month=1)
        end_text = match.group("end")
        end = None if _is_current_period_end(end_text) else _parse_year_month(end_text, default_month=12)
        if not start:
            continue
        periods.append(
            {
                "raw": raw,
                "start": _format_year_month(start),
                "end": "至今" if end is None else _format_year_month(end),
                "startIndex": _month_index(*start),
                "endIndex": None if end is None else _month_index(*end),
            }
        )
    return periods


def _looks_like_education_line(value: str) -> bool:
    markers = ["大学", "学院", "学校", "本科", "大专", "专科", "硕士", "博士", "专业", "教育经历"]
    return any(marker in value for marker in markers)


def _is_current_period_end(value: str) -> bool:
    return str(value or "").strip().lower() in {"至今", "现在", "目前", "今", "present", "在职"}


def _parse_year_month(value: str, default_month: int) -> tuple[int, int] | None:
    match = re.search(r"(?P<year>\d{4})(?:[./年](?P<month>\d{1,2}))?", str(value or ""))
    if not match:
        return None
    year = int(match.group("year"))
    month = int(match.group("month") or default_month)
    month = max(1, min(12, month))
    return year, month


def _format_year_month(value: tuple[int, int]) -> str:
    return f"{value[0]:04d}-{value[1]:02d}"


def _month_index(year: int, month: int) -> int:
    return year * 12 + month


def _conversation_has_resume_request_or_attachment(conversation: dict[str, Any]) -> bool:
    markers = [
        "\u6c42\u7b80\u5386",
        "\u8fd9\u662f\u6211\u7684\u7b80\u5386",
        "\u6211\u7684\u7b80\u5386",
        "\u7b80\u5386\u8bf7\u6c42\u5df2\u53d1\u9001",
        "\u5bf9\u65b9\u60f3\u53d1\u9001\u9644\u4ef6\u7b80\u5386",
        "\u70b9\u51fb\u9884\u89c8\u9644\u4ef6\u7b80\u5386",
        "\u9644\u4ef6\u7b80\u5386",
    ]
    return any(
        any(marker in (message.get("text") or "") for marker in markers)
        for message in conversation.get("messages", [])
    )


def _conversation_has_outbound_resume_request_text(conversation: dict[str, Any], message: str) -> bool:
    outbound_texts = [
        (item.get("text") or "").strip()
        for item in conversation.get("messages", [])
        if item.get("direction") == "outbound"
    ]
    return any(
        text == message.strip() or "\u65b9\u4fbf\u53d1\u4e00\u4efd\u7b80\u5386" in text
        for text in outbound_texts
    )


def _request_resume_new_greetings_stopped(
    job_text: str,
    message: str,
    criteria: str,
    dry_run: bool,
    inbox_scrolls: int,
    wait_after_unread: float,
    operation_delay_seconds: float,
    switch_status: dict[str, Any] | None,
    switch_job: dict[str, Any] | None,
    switch_unread: dict[str, Any] | None,
    inbox: dict[str, Any] | None,
    reason: str,
) -> dict[str, Any]:
    return _request_resume_new_greetings_payload(
        job_text, message, criteria, dry_run, inbox_scrolls, wait_after_unread, operation_delay_seconds,
        switch_status, switch_job, switch_unread, inbox, [], [], reason,
    )


def _request_resume_new_greetings_payload(
    job_text: str,
    message: str,
    criteria: str,
    dry_run: bool,
    inbox_scrolls: int,
    wait_after_unread: float,
    operation_delay_seconds: float,
    switch_status: dict[str, Any] | None,
    switch_job: dict[str, Any] | None,
    switch_unread: dict[str, Any] | None,
    inbox: dict[str, Any] | None,
    conversations: list[dict[str, Any]],
    results: list[dict[str, Any]],
    reason: str = "",
) -> dict[str, Any]:
    success_count = sum(1 for item in results if item.get("success"))
    skipped_count = sum(1 for item in results if item.get("skipped"))
    failed_count = sum(
        1 for item in results
        if not dry_run and not item.get("success") and not item.get("skipped")
    )
    return {
        "jobText": job_text,
        "message": message,
        "criteria": criteria,
        "dryRun": dry_run,
        "inboxScrolls": inbox_scrolls,
        "waitAfterUnread": wait_after_unread,
        "operationDelaySeconds": operation_delay_seconds,
        "switchStatusTab": switch_status,
        "switchJobFilter": switch_job,
        "switchUnreadFilter": switch_unread,
        "inbox": inbox,
        "snapshotCount": len(conversations),
        "processedCount": len(results),
        "successCount": success_count,
        "skippedCount": skipped_count,
        "failedCount": failed_count,
        "reason": reason,
        "results": results,
    }


def _reply_unread_with_knowledge(
    client: ChromeDebugClient,
    url_contains: str,
    job_text: str,
    max_count: int,
    inbox_scrolls: int,
    operation_delay_seconds: float,
    send: bool,
    knowledge_dir: Path | None = None,
    knowledge_file: Path | None = None,
) -> dict[str, Any]:
    dry_run = not send
    legacy_knowledge = _load_reply_knowledge_file(knowledge_file) if knowledge_file else None
    knowledge = load_knowledge_base(knowledge_dir) if knowledge_dir and knowledge_dir.exists() else None
    switch_job = client.switch_job_filter(job_text=job_text, url_contains=url_contains)
    if not switch_job.get("switched"):
        return _reply_unread_with_knowledge_payload(
            job_text=job_text,
            knowledge_file=knowledge_file,
            knowledge_dir=knowledge_dir,
            dry_run=dry_run,
            inbox_scrolls=inbox_scrolls,
            operation_delay_seconds=operation_delay_seconds,
            switch_job=switch_job,
            switch_unread=None,
            inbox=None,
            conversations=[],
            results=[],
            reason="job_filter_not_switched",
        )
    _sleep_operation_delay(operation_delay_seconds)

    switch_unread = client.switch_message_filter(label="\u672a\u8bfb", url_contains=url_contains)
    _sleep_operation_delay(operation_delay_seconds)
    inbox = client.capture_unread_inbox_scrolled(
        url_contains=url_contains,
        max_items=max_count,
        max_scrolls=inbox_scrolls,
    )
    conversations = [item for item in inbox.get("conversations", []) if item.get("dataId")][:max_count]
    results: list[dict[str, Any]] = []

    for item in conversations:
        open_result = None
        readiness = None
        fill_result = None
        send_result = None
        if dry_run:
            latest_message = str(item.get("summary") or "")
            latest_direction = "inbound" if latest_message else ""
            match = _match_unread_knowledge(
                latest_message=latest_message,
                latest_direction=latest_direction,
                job_title=str(item.get("jobTitle") or job_text),
                knowledge=knowledge,
                legacy_knowledge=legacy_knowledge,
            )
            skipped = match is None
            results.append(
                {
                    "target": item,
                    "openConversation": None,
                    "conversationReady": None,
                    "latestMessage": latest_message,
                    "latestDirection": latest_direction,
                    "matchedFaq": match.get("question") if match else "",
                    "matchedIntent": match.get("intent") if match else "",
                    "jobFolder": match.get("jobFolder") if match else "",
                    "knowledgeReason": match.get("reason") if match else "knowledge_not_matched",
                    "reply": match.get("answer") if match else "",
                    "fill": None,
                    "send": None,
                    "success": False,
                    "skipped": skipped,
                    "reason": "knowledge_not_matched" if skipped else "dry_run",
                }
            )
            _sleep_operation_delay(operation_delay_seconds)
            continue

        open_result = client.open_conversation(
            data_id=item.get("dataId"),
            candidate_name=item.get("candidateName"),
            url_contains=url_contains,
            max_scrolls=inbox_scrolls,
        )
        _sleep_operation_delay(operation_delay_seconds)
        readiness = client.wait_for_conversation_ready(
            expected_candidate_name=item.get("candidateName"),
            expected_job_title=item.get("jobTitle"),
            url_contains=url_contains,
        )
        if not readiness.get("ready"):
            results.append(
                _reply_unread_result(
                    target=item,
                    open_result=open_result,
                    readiness=readiness,
                    reason="conversation_not_ready",
                    skipped=True,
                )
            )
            continue

        conversation = readiness.get("conversation", {})
        latest = _latest_business_message(conversation)
        latest_text = str(latest.get("text") or "")
        latest_direction = str(latest.get("direction") or "")
        if latest_direction != "inbound":
            results.append(
                _reply_unread_result(
                    target=item,
                    open_result=open_result,
                    readiness=readiness,
                    latest_message=latest_text,
                    latest_direction=latest_direction,
                    reason="latest_message_not_inbound",
                    skipped=True,
                )
            )
            continue

        match = _match_unread_knowledge(
            latest_message=latest_text,
            latest_direction=latest_direction,
            job_title=str(conversation.get("jobTitle") or item.get("jobTitle") or job_text),
            knowledge=knowledge,
            legacy_knowledge=legacy_knowledge,
        )
        if match is None:
            results.append(
                _reply_unread_result(
                    target=item,
                    open_result=open_result,
                    readiness=readiness,
                    latest_message=latest_text,
                    latest_direction=latest_direction,
                    reason="knowledge_not_matched",
                    skipped=True,
                )
            )
            continue

        fill_result = client.fill_chat_input(match["answer"], url_contains=url_contains)
        _sleep_operation_delay(operation_delay_seconds)
        reason = ""
        if fill_result.get("filled"):
            send_result = client.send_current_message(url_contains=url_contains)
            _sleep_operation_delay(operation_delay_seconds)
            if not send_result.get("sent"):
                reason = str(send_result.get("reason") or "send_failed")
        else:
            reason = str(fill_result.get("reason") or "fill_failed")
        success = bool(send_result and send_result.get("sent"))
        results.append(
            _reply_unread_result(
                target=item,
                open_result=open_result,
                readiness=readiness,
                latest_message=latest_text,
                latest_direction=latest_direction,
                matched_question=match["question"],
                matched_intent=match.get("intent", ""),
                job_folder=match.get("jobFolder", ""),
                knowledge_reason=match.get("reason", "matched"),
                reply=match["answer"],
                fill=fill_result,
                send_result=send_result,
                success=success,
                skipped=False,
                reason="" if success else reason,
            )
        )

    return _reply_unread_with_knowledge_payload(
        job_text=job_text,
        knowledge_file=knowledge_file,
        knowledge_dir=knowledge_dir,
        dry_run=dry_run,
        inbox_scrolls=inbox_scrolls,
        operation_delay_seconds=operation_delay_seconds,
        switch_job=switch_job,
        switch_unread=switch_unread,
        inbox=inbox,
        conversations=conversations,
        results=results,
    )


def _load_reply_knowledge_file(path: Path) -> dict[str, Any]:
    content = path.read_text(encoding="utf-8")
    fallback = ""
    fallback_match = re.search(r"(?s)\[兜底回复\]\s*(.+?)(?=\n\[|$)", content)
    if fallback_match:
        fallback = fallback_match.group(1).strip()
    faq = _parse_reply_faq_text(content)
    return {"faq": faq, "fallback": fallback}


def _parse_reply_faq_text(content: str) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    question = ""
    answer_lines: list[str] = []

    def flush() -> None:
        nonlocal question, answer_lines
        if question and answer_lines:
            items.append({"question": question, "answer": "\n".join(answer_lines).strip()})
        question = ""
        answer_lines = []

    for raw_line in content.splitlines():
        line = raw_line.strip().lstrip("\ufeff")
        if not line:
            continue
        if re.match(r"^[Qq]\s*[:：]", line):
            flush()
            question = re.sub(r"^[Qq]\s*[:：]\s*", "", line).strip()
        elif re.match(r"^[Aa]\s*[:：]", line):
            answer_lines.append(re.sub(r"^[Aa]\s*[:：]\s*", "", line).strip())
        elif question:
            answer_lines.append(line)

    flush()
    return [item for item in items if item["question"] and item["answer"]]


def _match_reply_knowledge(message: str, faq_items: list[dict[str, str]]) -> dict[str, str] | None:
    query = normalize_text(message)
    if not query:
        return None
    query_tokens = _reply_match_tokens(query)
    if not query_tokens:
        return None
    for item in faq_items:
        question = normalize_text(item.get("question", ""))
        answer = item.get("answer", "")
        if not question or not answer:
            continue
        question_tokens = _reply_match_tokens(question)
        if not question_tokens:
            continue
        if question in query or query in question:
            return {"question": item.get("question", ""), "answer": answer}

    best_item: dict[str, str] | None = None
    best_score = 0
    for item in faq_items:
        question_tokens = _reply_match_tokens(normalize_text(item.get("question", "")))
        if not question_tokens:
            continue
        score = len(query_tokens & question_tokens)
        if score > best_score:
            best_score = score
            best_item = item
    if best_item and best_score >= 2:
        return {"question": best_item.get("question", ""), "answer": best_item.get("answer", "")}
    return None


def _match_unread_knowledge(
    latest_message: str,
    latest_direction: str,
    job_title: str,
    knowledge: dict[str, Any] | None,
    legacy_knowledge: dict[str, Any] | None,
) -> dict[str, str] | None:
    if latest_direction != "inbound":
        return None
    if legacy_knowledge is not None:
        match = _match_reply_knowledge(latest_message, legacy_knowledge["faq"])
        if not match:
            return None
        return {
            "question": match["question"],
            "answer": match["answer"],
            "intent": "legacy",
            "jobFolder": "",
            "reason": "matched",
        }
    if knowledge is None:
        return None

    qa_result = answer_question_from_knowledge_detailed(
        {
            "jobTitle": job_title,
            "messages": [{"direction": "inbound", "text": latest_message}],
        },
        knowledge,
    )
    if not qa_result.get("answer"):
        return None
    matches = qa_result.get("matches", [])
    return {
        "question": " / ".join(str(item.get("question", "")) for item in matches if item.get("question")),
        "answer": str(qa_result.get("answer") or ""),
        "intent": ",".join(str(item.get("intent", "")) for item in matches if item.get("intent")),
        "jobFolder": str(qa_result.get("jobFolder") or ""),
        "reason": str(qa_result.get("reason") or "matched"),
    }


def _reply_match_tokens(text: str) -> set[str]:
    chunks = re.findall(r"[\u4e00-\u9fff]{2,}|[a-zA-Z0-9]+", text)
    tokens: set[str] = set()
    for chunk in chunks:
        if re.fullmatch(r"[\u4e00-\u9fff]+", chunk):
            tokens.update(chunk[index:index + 2] for index in range(max(0, len(chunk) - 1)))
        else:
            tokens.add(chunk.lower())
    return tokens


def _latest_business_message(conversation: dict[str, Any]) -> dict[str, Any]:
    for message in reversed(conversation.get("messages", [])):
        if message.get("direction") in {"inbound", "outbound"}:
            return message
    return {}


def _reply_unread_result(
    target: dict[str, Any],
    open_result: dict[str, Any] | None,
    readiness: dict[str, Any] | None,
    latest_message: str = "",
    latest_direction: str = "",
    matched_question: str = "",
    matched_intent: str = "",
    job_folder: str = "",
    knowledge_reason: str = "",
    reply: str = "",
    fill: dict[str, Any] | None = None,
    send_result: dict[str, Any] | None = None,
    success: bool = False,
    skipped: bool = False,
    reason: str = "",
) -> dict[str, Any]:
    return {
        "target": target,
        "openConversation": open_result,
        "conversationReady": readiness,
        "latestMessage": latest_message,
        "latestDirection": latest_direction,
        "matchedFaq": matched_question,
        "matchedIntent": matched_intent,
        "jobFolder": job_folder,
        "knowledgeReason": knowledge_reason,
        "reply": reply,
        "fill": fill,
        "send": send_result,
        "success": success,
        "skipped": skipped,
        "reason": reason,
    }


def _reply_unread_with_knowledge_payload(
    job_text: str,
    knowledge_file: Path | None,
    knowledge_dir: Path | None,
    dry_run: bool,
    inbox_scrolls: int,
    operation_delay_seconds: float,
    switch_job: dict[str, Any] | None,
    switch_unread: dict[str, Any] | None,
    inbox: dict[str, Any] | None,
    conversations: list[dict[str, Any]],
    results: list[dict[str, Any]],
    reason: str = "",
) -> dict[str, Any]:
    success_count = sum(1 for item in results if item.get("success"))
    skipped_count = sum(1 for item in results if item.get("skipped"))
    failed_count = sum(
        1 for item in results
        if not dry_run and not item.get("success") and not item.get("skipped")
    )
    return {
        "jobText": job_text,
        "knowledgeFile": str(knowledge_file) if knowledge_file else "",
        "knowledgeDir": str(knowledge_dir) if knowledge_dir else "",
        "knowledgeMode": "legacy_file" if knowledge_file else "routed_dir",
        "dryRun": dry_run,
        "inboxScrolls": inbox_scrolls,
        "operationDelaySeconds": operation_delay_seconds,
        "switchJobFilter": switch_job,
        "switchUnreadFilter": switch_unread,
        "inbox": inbox,
        "snapshotCount": len(conversations),
        "processedCount": len(results),
        "successCount": success_count,
        "skippedCount": skipped_count,
        "failedCount": failed_count,
        "reason": reason,
        "results": results,
    }


def _send_unread_message(
    client: ChromeDebugClient,
    url_contains: str,
    job_text: str,
    message: str,
    max_count: int,
    inbox_scrolls: int,
    delay_seconds: float,
    dry_run: bool,
) -> dict[str, Any]:
    switch_job = client.switch_job_filter(job_text=job_text, url_contains=url_contains)
    switch_unread = client.switch_message_filter(label="\u672a\u8bfb", url_contains=url_contains)
    inbox = client.capture_unread_inbox_scrolled(
        url_contains=url_contains,
        max_items=max_count,
        max_scrolls=inbox_scrolls,
    )
    conversations = [
        item for item in inbox.get("conversations", [])
        if item.get("dataId")
    ][:max_count]
    switch_all = None
    results: list[dict[str, Any]] = []

    if conversations and not dry_run:
        switch_all = client.switch_message_filter(label="\u5168\u90e8", url_contains=url_contains)

    for item in conversations:
        if dry_run:
            results.append(
                {
                    "target": item,
                    "openConversation": None,
                    "conversationReady": None,
                    "fill": None,
                    "send": None,
                    "closeAttachmentPreview": None,
                    "sent": False,
                    "skipped": True,
                    "reason": "dry_run",
                }
            )
            continue

        open_result = client.open_conversation(
            data_id=item.get("dataId"),
            candidate_name=item.get("candidateName"),
            url_contains=url_contains,
            max_scrolls=inbox_scrolls,
        )
        readiness = client.wait_for_conversation_ready(
            expected_candidate_name=item.get("candidateName"),
            expected_job_title=item.get("jobTitle"),
            url_contains=url_contains,
        )
        if not readiness.get("ready"):
            close_result = client.close_attachment_preview(url_contains=url_contains)
            results.append(
                {
                    "target": item,
                    "openConversation": open_result,
                    "conversationReady": readiness,
                    "fill": None,
                    "send": None,
                    "closeAttachmentPreview": close_result,
                    "sent": False,
                    "skipped": True,
                    "reason": "conversation_not_ready",
                }
            )
            continue

        fill_result = client.fill_chat_input(message, url_contains=url_contains)
        send_result = None
        reason = ""
        if fill_result.get("filled"):
            send_result = client.send_current_message(url_contains=url_contains)
            if not send_result.get("sent"):
                reason = str(send_result.get("reason") or "send_failed")
        else:
            reason = str(fill_result.get("reason") or "fill_failed")
        close_result = client.close_attachment_preview(url_contains=url_contains)
        sent = bool(send_result and send_result.get("sent"))
        results.append(
            {
                "target": item,
                "openConversation": open_result,
                "conversationReady": readiness,
                "fill": fill_result,
                "send": send_result,
                "closeAttachmentPreview": close_result,
                "sent": sent,
                "skipped": False,
                "reason": reason,
            }
        )
        if delay_seconds > 0:
            time.sleep(delay_seconds)

    sent_count = sum(1 for item in results if item.get("sent"))
    failed_count = sum(
        1 for item in results
        if not dry_run and not item.get("sent")
    )
    return {
        "jobText": job_text,
        "message": message,
        "dryRun": dry_run,
        "inboxScrolls": inbox_scrolls,
        "switchJobFilter": switch_job,
        "switchUnreadFilter": switch_unread,
        "inbox": inbox,
        "snapshotCount": len(conversations),
        "switchOpenFilter": switch_all,
        "sentCount": sent_count,
        "failedCount": failed_count,
        "results": results,
    }


def _process_conversation_item(
    client: ChromeDebugClient,
    item: dict[str, Any],
    url_contains: str,
    downloads_dir: Path,
    knowledge_dir: Path,
    send: bool,
    delay_seconds: float,
) -> dict[str, Any]:
    open_result = client.open_conversation(
        data_id=item.get("dataId"),
        candidate_name=item.get("candidateName"),
        url_contains=url_contains,
    )
    readiness = client.wait_for_conversation_ready(
        expected_candidate_name=item.get("candidateName"),
        expected_job_title=item.get("jobTitle"),
        url_contains=url_contains,
    )
    close_result = None

    if not readiness.get("ready"):
        close_result = client.close_attachment_preview(url_contains=url_contains)
        return {
            "target": item,
            "openConversation": open_result,
            "conversationReady": readiness,
            "prepareReply": None,
            "send": None,
            "closeAttachmentPreview": close_result,
            "skipped": True,
            "reason": "conversation_not_ready",
        }

    if not _has_new_inbound_reply(readiness.get("conversation", {})):
        close_result = client.close_attachment_preview(url_contains=url_contains)
        return {
            "target": item,
            "openConversation": open_result,
            "conversationReady": readiness,
            "prepareReply": None,
            "send": None,
            "closeAttachmentPreview": close_result,
            "skipped": True,
            "reason": "no_new_reply",
        }

    prepare_result = _prepare_reply(
        client=client,
        url_contains=url_contains,
        downloads_dir=downloads_dir,
        knowledge_dir=knowledge_dir,
    )
    send_result = None
    auto_send_reason = _auto_send_reason(prepare_result)
    should_send = send and auto_send_reason == ""
    if should_send:
        send_result = client.send_current_message(url_contains=url_contains)
    close_result = client.close_attachment_preview(url_contains=url_contains)
    if delay_seconds > 0:
        time.sleep(delay_seconds)
    return {
        "target": item,
        "openConversation": open_result,
        "conversationReady": readiness,
        "prepareReply": prepare_result,
        "send": send_result,
        "closeAttachmentPreview": close_result,
        "skipped": False,
        "reason": "" if should_send or not send else auto_send_reason,
    }


def _is_today_time(value: str) -> bool:
    value = value.strip()
    return ":" in value


def _resolve_delay_seconds(min_delay_seconds: float, max_delay_seconds: float) -> float:
    lower = max(0.0, min(min_delay_seconds, max_delay_seconds))
    upper = max(0.0, max(min_delay_seconds, max_delay_seconds))
    if upper == lower:
        return lower
    return random.uniform(lower, upper)


def _has_new_inbound_reply(conversation: dict[str, Any]) -> bool:
    messages = conversation.get("messages", [])
    business_messages = [
        message for message in messages
        if message.get("direction") in {"inbound", "outbound"}
    ]
    if not business_messages:
        return False
    return business_messages[-1].get("direction") == "inbound"


def _auto_send_reason(prepare_result: dict[str, Any]) -> str:
    draft = prepare_result.get("draft", {})
    body = (draft.get("body") or "").strip()
    score = prepare_result.get("score", {})
    conversation = prepare_result.get("conversation", {})
    messages = conversation.get("messages", [])
    resume_download = prepare_result.get("resumeDownload")
    requires_resume_download = any(
        "对方想发送附件简历给您，您是否同意" in (message.get("text") or "")
        or "点击预览附件简历" in (message.get("text") or "")
        or "附件简历" in (message.get("text") or "")
        for message in messages
    )
    outbound_texts = [
        (message.get("text") or "").strip()
        for message in messages
        if message.get("direction") == "outbound"
    ]
    latest_outbound = outbound_texts[-1] if outbound_texts else ""

    if not body:
        return "empty_draft"
    if not messages:
        return "empty_conversation"
    if requires_resume_download and not (resume_download and resume_download.get("downloaded")):
        return "resume_not_downloaded"
    if body == "您好，已收到您的消息。我先看一下当前沟通内容，稍后再给您更准确的回复。":
        return "low_confidence_reply"
    if score.get("recommended_action") == "review":
        return "review_only"
    if body in outbound_texts or body == latest_outbound:
        return "duplicate_reply"
    return ""


def _download_resume_with_retry(
    client: ChromeDebugClient,
    url_contains: str,
    downloads_dir: Path,
    attempts: int = 3,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    latest_preview: dict[str, Any] | None = None
    latest_download: dict[str, Any] | None = None
    for _ in range(attempts):
        latest_preview = client.capture_attachment_preview_link(url_contains=url_contains)
        latest_download = client.download_resume(
            url_contains=url_contains,
            downloads_dir=downloads_dir,
        )
        if latest_download.get("downloaded"):
            return latest_preview, latest_download
        time.sleep(1)
    return latest_preview, latest_download


def _parse_filter_args(values: list[str]) -> dict[str, str]:
    filters: dict[str, str] = {}
    for value in values:
        if "=" not in value:
            filters[value] = value
            continue
        key, raw = value.split("=", 1)
        key = key.strip()
        raw = raw.strip()
        if key in filters and raw:
            filters[key] = f"{filters[key]},{raw}"
        else:
            filters[key] = raw
    return {key: raw for key, raw in filters.items() if key}


def _process_recommendations(
    client: ChromeDebugClient,
    url_contains: str,
    job_title: str,
    city: str,
    filters: dict[str, str],
    knowledge_dir: Path,
    max_count: int,
    min_score: int,
    send: bool,
) -> dict[str, Any]:
    open_result = client.open_recommend_page(url_contains=url_contains)
    search_result = client.apply_recommend_filters(
        job_title=job_title,
        city=city,
        filters=filters,
        url_contains=url_contains,
    )
    candidates_payload = client.capture_recommend_candidates(
        max_count=max_count,
        url_contains=url_contains,
    )
    processed: list[dict[str, Any]] = []
    processed_keys: set[str] = set()

    for item in candidates_payload.get("candidates", [])[:max_count]:
        candidate_key = str(item.get("candidateId") or item.get("cardIndex") or item.get("candidateName") or "")
        if candidate_key in processed_keys:
            processed.append(
                {
                    "candidate": item,
                    "skipped": True,
                    "reason": "duplicate_candidate",
                    "score": None,
                    "greet": None,
                }
            )
            continue
        processed_keys.add(candidate_key)

        try:
            open_candidate = client.open_recommend_candidate(
                candidate_id=str(item.get("candidateId") or ""),
                card_index=int(item.get("cardIndex") or 0),
                url_contains=url_contains,
            )
            if not open_candidate.get("opened"):
                processed.append(
                    {
                        "candidate": item,
                        "openCandidate": open_candidate,
                        "resume": None,
                        "score": None,
                        "greet": None,
                        "skipped": True,
                        "reason": "candidate_not_opened",
                    }
                )
                continue

            time.sleep(0.5)
            resume = client.capture_candidate_resume(url_contains=url_contains)
            if not resume.get("rawTextPreview") and not resume.get("name"):
                processed.append(
                    {
                        "candidate": item,
                        "openCandidate": open_candidate,
                        "resume": resume,
                        "score": None,
                        "greet": None,
                        "skipped": True,
                        "reason": "resume_not_captured",
                    }
                )
                continue

            score = score_recommendation(
                resume,
                job_title=job_title,
                knowledge_dir=knowledge_dir,
                min_score=min_score,
            )
            should_greet = score.get("decision") == "greet" and score.get("match_score", 0) >= min_score
            greet_result = None
            if should_greet and score.get("greeting"):
                greet_result = client.greet_candidate(
                    message=score["greeting"],
                    send=send,
                    url_contains=url_contains,
                )

            processed.append(
                {
                    "candidate": item,
                    "openCandidate": open_candidate,
                    "resume": resume,
                    "score": score,
                    "greet": greet_result,
                    "skipped": not should_greet,
                    "reason": "" if should_greet else score.get("decision", "skip"),
                }
            )
        except Exception as exc:
            processed.append(
                {
                    "candidate": item,
                    "resume": None,
                    "score": None,
                    "greet": None,
                    "skipped": True,
                    "reason": "error",
                    "error": str(exc),
                }
            )

    return {
        "openRecommendPage": open_result,
        "search": search_result,
        "candidates": candidates_payload,
        "processedCount": len(processed),
        "processed": processed,
        "sent": send,
        "minScore": min_score,
    }


if __name__ == "__main__":
    raise SystemExit(main())
